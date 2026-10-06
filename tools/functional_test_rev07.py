"""Functional test for the revision 07 TAB workbook (the cells and formulas tools/build_rev07.py adds or moves).

    python3 tools/functional_test_rev07.py "07 - a2b_Blank_TAB_Workbook <date>.xlsm" ["06 - a2b_Blank_TAB_Workbook <date>.xlsm"]

Copies of the workbook are filled with openpyxl, recalculated with LibreOffice (soffice headless) and checked:

  0. blank template: no error cells
  1. unit blocks: drive rows show the sheave / pulley from the data entry; Estimated BHP = HP x avg amps / FLA (nameplate
     HP, else scheduled; corrected FLA, else nameplate FLA; an N/A nameplate HP falls back)
  2. MAU page 2: PSP initial / final CFM, filter grid initial / final CFM and totals, intake screens, profile pressure,
     Method Total for every method (final, else initial), page-1 total airflow, page-2 outlets in the totals
  3. Fans: the measured-at-hood line in the fan totals and its % of design
  4. Building Balance: an Excl. unit is marked, left out of the totals (both sides)
  5. Traverses: a flat oval (rectangle + ends, area-weighted), and a rectangular and a round traverse with the same
     readings give the same values as revision 06
  6. N/A notations in the new inputs: no error cells
"""
import glob
import math
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = sys.argv[1]
REV06 = sys.argv[2] if len(sys.argv) > 2 else sorted(glob.glob(os.path.join(ROOT, "06 - a2b_Blank_TAB_Workbook *.xlsm")))[-1]
tmpdir = tempfile.mkdtemp()
ERRS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "Err:")
results = []


def recalc(path):
    outdir = os.path.join(tmpdir, "conv")
    os.makedirs(outdir, exist_ok=True)
    r = subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", outdir, path],
                       capture_output=True, text=True, timeout=900)
    conv = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".xlsx")
    if not os.path.exists(conv):
        raise SystemExit(f"soffice could not recalculate {path}: {r.stdout} {r.stderr}")
    shutil.move(conv, path)


def prepared(src, name, *fills):
    path = os.path.join(tmpdir, name + ".xlsm")
    shutil.copy(src, path)
    if fills:
        wb = load_workbook(path, keep_vba=True)
        for f in fills:
            f(wb)
        wb.save(path)
    recalc(path)
    return load_workbook(path, data_only=True)


def errors(wb):
    return [f"{ws.title}!{c.coordinate}={c.value}" for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith(ERRS)]


def check(section, name, got, exp, tol=1e-6):
    ok = (got == exp) or (isinstance(got, (int, float)) and isinstance(exp, (int, float))
                          and not isinstance(got, bool) and abs(got - exp) < tol)
    results.append((section, name, got, exp, ok))


# ----------------------------------------------------------------------------------------- fills
P1 = 4          # first unit block
Q1 = P1 + 52    # its continuation page
EDE = "{Equipment Data Entry}"


def fill_units(wb):
    e = wb[EDE]
    for row, des in ((7, "RTU-1"), (52, "MUA-1"), (81, "EF-1")):
        e[f"B{row}"] = des
        e[f"G{row}"] = 5           # scheduled HP
        e[f"N{row}"] = 460         # rated voltage
        e[f"O{row}"] = "3-phase"
        e[f"J{row}"] = "2VP60"     # motor sheave
        e[f"K{row}"] = "BK90"      # fan pulley
    for name in ("RTUs", "MAUs", "Fans"):
        ws = wb[name]
        ws["E17"] = 20             # FLA
        for c, v in zip("EFG", (460, 460, 460)):
            ws[f"{c}19"] = v
        for c, v in zip("EFG", (10, 11, 12)):
            ws[f"{c}20"] = v
        ws["M16"] = "1-3/8"
        ws["M17"] = "1-7/16"
    wb["RTUs"]["E21"] = 7.5        # nameplate HP
    wb["MAUs"]["E21"] = "N/A"      # falls back to the scheduled HP
    # Fans: corrected FLA (volts 440 -> 460/440*20) and no nameplate HP
    for c in "EFG":
        wb["Fans"][f"{c}19"] = 440


def fill_mau(method):
    def f(wb):
        ws = wb["MAUs"]
        Q = Q1
        ws[f"D{Q + 3}"] = 48
        ws[f"G{Q + 3}"] = 10
        ws[f"J{Q + 3}"] = 0
        for r in (Q + 4, Q + 5):
            for c in "DEFGHIJKLM":
                ws[f"{c}{r}"] = 400
        if method != "PSP-initial":
            for r in (Q + 6, Q + 7):
                for c in "DEFGHIJKLM":
                    ws[f"{c}{r}"] = 500
        ws[f"C{Q + 11}"] = '12" x 12"'
        ws[f"D{Q + 11}"] = '12" x 24"'
        ws[f"C{Q + 12}"], ws[f"D{Q + 12}"] = 300, 310
        ws[f"C{Q + 13}"], ws[f"D{Q + 13}"] = 350, 360
        ws[f"D{Q + 19}"] = 1
        ws[f"H{Q + 19}"] = 0.15
        for i, (ak, iv, fv) in enumerate(((4, 200, 250), (5, 220, 260))):
            r = Q + 23 + i
            ws[f"B{r}"] = f"S-{i + 1}"
            ws[f"C{r}"] = "48x24"
            ws[f"F{r}"] = ak
            ws[f"I{r}"] = iv
            if method != "Intake-initial":
                ws[f"K{r}"] = fv
        ws[f"E{Q + 30}"] = method.split("-")[0]
        # one outlet on page 1, one on page 2
        for r, d in ((P1 + 29, 1000), (Q + 35, 500)):
            ws[f"B{r}"], ws[f"F{r}"], ws[f"H{r}"], ws[f"K{r}"] = "S", 1, d, d * 0.9
    return f


def fill_fan_hood(wb):
    ws = wb["Fans"]
    r = P1 + 44
    ws[f"C{r}"] = "Measured at hood H-1, H-2"
    ws[f"H{r}"], ws[f"J{r}"], ws[f"L{r}"] = 1000, 900, 950
    ws[f"B{P1 + 29}"], ws[f"F{P1 + 29}"], ws[f"H{P1 + 29}"], ws[f"K{P1 + 29}"] = "G-1", 1, 200, 210


def fill_balance(wb):
    e = wb[EDE]
    for i in range(3):
        e[f"B{7 + i}"] = f"RTU-{i + 1}"
        e[f"B{81 + i}"] = f"EF-{i + 1}"
    for i in range(3):
        rtu = wb["RTUs"]
        P = 4 + 104 * i
        rtu[f"B{P + 47}"], rtu[f"F{P + 47}"], rtu[f"H{P + 47}"], rtu[f"K{P + 47}"] = "OA", 1, 100 * (i + 1), 110 * (i + 1)
        fan = wb["Fans"]
        fan[f"B{P + 29}"], fan[f"F{P + 29}"], fan[f"H{P + 29}"], fan[f"K{P + 29}"] = "G", 1, 50 * (i + 1), 60 * (i + 1)
    bb = wb["Building Balance"]
    bb["P8"] = "Excl."
    bb["Q9"] = "Excl."


def fill_traverses(flat=True):
    def f(wb):
        ws = wb["Traverses"]
        # T-1 flat oval 30 x 12: 3 x 3 rectangle readings 1000, 8 end points 800
        if flat:
            ws["D9"], ws["G9"], ws["I9"] = "Flat Oval", 30, 12
            vals = [1000] * 9 + [800] * 8
            for k, v in enumerate(vals):
                ws[f"{'PQRSTUVW'[k // 10]}{11 + k % 10}"] = v
        # T-2 rectangular 24 x 12, T-3 round 14
        for T, shape, w, h in ((20, "Rectangular", 24, 12), (35, "Round", 14, None)):
            ws[f"D{T + 4}"], ws[f"G{T + 4}"] = shape, w
            if h:
                ws[f"I{T + 4}"] = h
            for k in range(20):
                ws[f"{'PQRSTUVW'[k // 10]}{T + 6 + k % 10}"] = 900 + 10 * k
    return f


def fill_na(wb):
    for name in ("RTUs", "MAUs", "ERVs", "Fans"):
        ws = wb[name]
        for ref in ("M16", "M17", "E21", "E20", "F20", "G20"):
            ws[ref] = "N/A"
    ws = wb["MAUs"]
    for ref in (f"D{Q1 + 4}", f"D{Q1 + 6}", f"C{Q1 + 12}", f"C{Q1 + 13}", f"F{Q1 + 23}", f"I{Q1 + 23}", f"K{Q1 + 23}"):
        ws[ref] = "N/A"
    ws[f"C{Q1 + 11}"] = '12" x 12"'
    ws[f"E{Q1 + 30}"] = "Intake"
    tr = wb["Traverses"]
    tr["D9"], tr["G9"], tr["I9"] = "Flat Oval", "N/A", 12
    tr["D24"], tr["G24"], tr["I24"] = "Flat Oval", 30, "N/A"
    wb["Fans"][f"H{P1 + 44}"] = "N/A"
    wb["Building Balance"]["P7"] = "N/A"


# ----------------------------------------------------------------------------------------- checks
ONLY = set(os.environ.get("ONLY", "0123456"))


def run():
    sec = lambda k: k in ONLY  # noqa: E731
    if sec("0"):
        _run0()
    if sec("1"):
        _run1()
    if sec("2"):
        _run2()
    if sec("3"):
        _run3()
    if sec("4"):
        _run4()
    if sec("5"):
        _run5()
    if sec("6"):
        _run6()


def _run0():
    wb = prepared(SRC, "blank")
    errs = errors(wb)
    check("0", "blank template: error cells", len(errs), 0)


def _run1():
    wb = prepared(SRC, "units", fill_units)
    for name, hp, fla in (("RTUs", 7.5, 20), ("MAUs", 5, 20), ("Fans", 5, 460 / 440 * 20)):
        ws = wb[name]
        check("1", f"{name} motor sheave K16", ws["K16"].value, "2VP60")
        check("1", f"{name} fan pulley K17", ws["K17"].value, "BK90")
        check("1", f"{name} motor bore M16", ws["M16"].value, "1-3/8")
        check("1", f"{name} Estimated BHP G18", ws["G18"].value, hp * 11 / fla)
    check("1", "unit blocks: error cells", len(errors(wb)), 0)


def _run2():
    sizes = {'12" x 12"': 0.69, '12" x 24"': 1.52}
    psp_i = 400 * (48 - 2) * 10 * 0.88 / 144
    psp_f = 500 * (48 - 2) * 10 * 0.88 / 144
    fg_i = 300 * 0.69 * 1.35 + 310 * 1.52 * 1.35
    fg_f = 350 * 0.69 * 1.35 + 360 * 1.52 * 1.35
    expect = {"PSP": psp_f, "PSP-initial": psp_i, "Filter Grid": fg_f, "Profile Pressure": 697.15,
              "Intake": 4 * 250 + 5 * 260, "Intake-initial": 4 * 200 + 5 * 220, "Outlets": 900 + 450}
    for method, total in expect.items():
        wb = prepared(SRC, f"mau-{method}", fill_mau(method))
        ws = wb["MAUs"]
        Q = Q1
        if method == "PSP":
            check("2", "PSP CFM initial", ws[f"E{Q + 8}"].value, psp_i)
            check("2", "PSP CFM final", ws[f"K{Q + 8}"].value, psp_f)
            check("2", "filter CFM initial C", ws[f"C{Q + 14}"].value, 300 * sizes['12" x 12"'] * 1.35)
            check("2", "filter CFM final D", ws[f"D{Q + 15}"].value, 360 * sizes['12" x 24"'] * 1.35)
            check("2", "filter total initial", ws[f"E{Q + 16}"].value, fg_i)
            check("2", "filter total final", ws[f"K{Q + 16}"].value, fg_f)
            check("2", "intake CFM row 1 final", ws[f"L{Q + 23}"].value, 1000)
            check("2", "intake total initial", ws[f"J{Q + 27}"].value, 1900)
            check("2", "intake total final", ws[f"L{Q + 27}"].value, 2300)
            check("2", "profile CFM", ws[f"K{Q + 19}"].value, 697.15)
            check("2", "page-2 outlet subtotal (final)", ws[f"L{Q + 47}"].value, 450)
            check("2", "page-1 total design (both pages)", ws[f"H{P1 + 45}"].value, 1500)
            check("2", "design total airflow K9", ws["K9"].value, 1500)
        check("2", f"Method Total ({method})", ws[f"E{Q + 31}"].value if method != "Outlets" else ws["L9"].value, total, 1e-6)
        check("2", f"Total Airflow actual L9 ({method})", ws["L9"].value, total, 1e-6)
        check("2", f"{method}: error cells", len(errors(wb)), 0)


def _run3():
    wb = prepared(SRC, "fan-hood", fill_fan_hood)
    ws = wb["Fans"]
    r = P1 + 44
    check("3", "hood line % (final / design)", ws[f"M{r}"].value, 0.95)
    check("3", "fan total design H49", ws[f"H{P1 + 45}"].value, 1200)
    check("3", "fan total final L49", ws[f"L{P1 + 45}"].value, 1160)
    check("3", "fan design K9 / actual L9", (ws["K9"].value, ws["L9"].value), (1200, 1160))
    check("3", "Building Balance EF-1 actual", wb["Building Balance"]["K7"].value, 1160)


def _run4():
    wb = prepared(SRC, "balance", fill_balance)
    bb = wb["Building Balance"]
    check("4", "RTU-2 marked (its %)", (bb["B8"].value, bb["G8"].value), ("RTU-2", "Excl."))
    check("4", "RTU-1 not marked", (bb["B7"].value, bb["G7"].value), ("RTU-1", 1.1))
    check("4", "OA design total without RTU-2", bb["C87"].value, 100 + 300)
    check("4", "OA actual total without RTU-2", bb["E87"].value, 110 + 330)
    check("4", "EF-3 marked (its %)", (bb["H9"].value, bb["M9"].value), ("EF-3", "Excl."))
    check("4", "exhaust design total without EF-3", bb["I87"].value, 50 + 100)
    check("4", "exhaust actual total without EF-3", bb["K87"].value, 60 + 120)
    check("4", "design building balance", bb["H89"].value, 400 - 150)


def _run5():
    wb7 = prepared(SRC, "trav7", fill_traverses(True))
    wb6 = prepared(REV06, "trav6", fill_traverses(False))
    t = wb7["Traverses"]
    ar, ac = 18 * 12 / 144, math.pi * 36 / 144
    vel = round((1000 * ar + 800 * ac) / (ar + ac))
    check("5", "flat oval Ak H7", t["H7"].value, round(ar + ac, 3))
    check("5", "flat oval size F7", t["F7"].value, '30" x 12" oval')
    check("5", "flat oval points M9", t["M9"].value, "3x3+8 ends")
    check("5", "flat oval end points O9", t["O9"].value, 8)
    check("5", "flat oval Final VEL L7", t["L7"].value, vel)
    check("5", "flat oval CFM M7", t["M7"].value, round(vel * round(ar + ac, 3)))
    check("5", "flat oval grid row 1", [t[f"{c}11"].value for c in "DEFG"], [1000, 1000, 1000, None])
    check("5", "flat oval ends row 4", [t[f"{c}14"].value for c in "DEFGHIJKL"], [800] * 8 + [None])
    check("5", "flat oval row labels", [t[f"B{r}"].value for r in (11, 13, 14, 15)], [2, 10, "Ends", None])
    check("5", "flat oval positions", [t[f"{c}10"].value for c in "DEFG"], [3, 9, 15, None])
    t6 = wb6["Traverses"]
    for T in (20, 35):
        for ref in (f"F{T + 2}", f"H{T + 2}", f"L{T + 2}", f"M{T + 2}", f"M{T + 4}", f"N{T + 4}", f"N{T + 5}", f"N{T + 6}"):
            check("5", f"same as rev 06: {ref}", t[ref].value, t6[ref].value)
        grid = [t[f"{c}{r}"].value for r in range(T + 5, T + 14) for c in "BDEFGHIJKLM"]
        grid6 = [t6[f"{c}{r}"].value for r in range(T + 5, T + 14) for c in "BDEFGHIJKLM"]
        check("5", f"same as rev 06: grid of block at row {T}", grid, grid6)
    check("5", "traverses: error cells", len(errors(wb7)), 0)
    if errors(wb7):
        print("  errors:", errors(wb7)[:10])


def _run6():
    wb = prepared(SRC, "na", fill_units, fill_na)
    errs = errors(wb)
    check("6", "N/A notations in new inputs: error cells", len(errs), 0)
    if errs:
        print("  errors:", errs[:20])
    check("6", "RTU BHP with N/A amps legs (E:G all N/A) is blank", wb["RTUs"]["G18"].value, None)


if __name__ == "__main__":
    run()
    bad = [r for r in results if not r[4]]
    for sec, name, got, exp, ok in results:
        print(f"{'PASS' if ok else 'FAIL'} [{sec}] {name}: {got!r}" + ("" if ok else f" (expected {exp!r})"))
    print(f"\n{len(results) - len(bad)} / {len(results)} checks passed")
    sys.exit(1 if bad else 0)
