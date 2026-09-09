"""Build the revised a2b TAB workbook from the 4-16-26 original.

    python3 tools/build_workbook.py            # writes 01 - a2b_Blank_TAB_Workbook <date>.xlsm

Every change is a function below so the diff against the original is auditable.
The original file is never modified.
"""
import copy
import datetime as dt
import os
import re
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break, RowBreak

sys.path.insert(0, os.path.dirname(__file__))
import xlsm_parts  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORIG = os.path.join(ROOT, "00 - a2b_Blank_TAB_Workbook 4-16-26.xlsm")
STAMP = os.environ.get("TAB_BUILD_DATE", dt.date.today().strftime("%-m-%-d-%y"))
OUT = os.path.join(ROOT, f"01 - a2b_Blank_TAB_Workbook {STAMP}.xlsm")

EDE = "'{Equipment Data Entry}'"
DD = "'{Dropdowns}'"
LOG = []


def log(msg):
    LOG.append(msg)
    print("  -", msg)


def anchors(ws, label="System", col="B"):
    return [c.row for c in ws[col] if c.value == label]


def f(ws, coord):
    return ws[coord].value


def setf(ws, coord, value, note=None):
    old = ws[coord].value
    if old != value:
        ws[coord].value = value
        if note:
            log(f"{ws.title}!{coord}: {note}")


def copy_style(src, dst):
    dst.font = copy.copy(src.font)
    dst.border = copy.copy(src.border)
    dst.fill = copy.copy(src.fill)
    dst.number_format = src.number_format
    dst.alignment = copy.copy(src.alignment)
    dst.protection = copy.copy(src.protection)


# --------------------------------------------------------------------------- #
# 1. Link defects
# --------------------------------------------------------------------------- #
def fix_fan_data_totals(wb):
    d = wb["Fan Data (EFs, TFs, etc.)"]; a = wb["Fan Airflow"]
    da, aa = anchors(d), anchors(a)
    n = 0
    for i, anc in enumerate(da):
        for r in range(anc, anc + 23):
            if d.cell(r, 9).value == "Total Airflow":
                t = aa[i] + 19
                for col, letter in ((11, "H"), (12, "L")):
                    want = f"='Fan Airflow'!{letter}{t}"
                    if d.cell(r, col).value != want:
                        d.cell(r, col).value = want; n += 1
    log(f"Fan Data: re-pointed {n} Total Airflow links to the matching Fan Airflow block")


def fix_mau_airflow_ak(wb):
    ws = wb["MAU Airflow"]; n = 0
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "$F$9" in c.value and c.row != 9:
                c.value = c.value.replace("$F$9", f"$F{c.row}"); n += 1
    log(f"MAU Airflow: replaced absolute $F$9 Ak reference in {n} CFM cells")


def fix_rtu_data_blocks(wb):
    ws = wb["RTU Data"]
    for anc in anchors(ws):
        tot, oa, ra, esp = anc + 5, anc + 6, anc + 7, anc + 8
        # Return airflow = total - OA (design and actual)
        for col in ("K", "L"):
            want = f'=IF({col}{tot}="","",{col}{tot}-N({col}{oa}))'
            cur = ws[f"{col}{ra}"].value
            if cur != want:
                ws[f"{col}{ra}"].value = want
                if cur is None or "RTU Airflow" in str(cur) or f"{col}{tot}" not in str(cur):
                    log(f"RTU Data!{col}{ra}: return airflow formula restored (was {cur!r})")
        # Unit ESP actual = fan exiting SP profile cell F(anc+21)
        want = f"=F{anc + 21}"
        cur = ws[f"L{esp}"].value
        if cur != want:
            ws[f"L{esp}"].value = want; log(f"RTU Data!L{esp}: Unit ESP actual restored (was {cur!r})")
        # stray links in header rows
        for r in (anc + 4, anc + 10):
            for col in ("K", "L"):
                v = ws[f"{col}{r}"].value
                if isinstance(v, str) and "RTU Airflow" in v:
                    ws[f"{col}{r}"].value = None; log(f"RTU Data!{col}{r}: cleared stray link {v}")
    setf(ws, "C40", f"={EDE}!O8", "phase now links to RTU-2 (was RTU-1's O7)")


# --------------------------------------------------------------------------- #
# 2. Formula hygiene
# --------------------------------------------------------------------------- #
RE_IF0 = re.compile(r"^=IF\((?P<ref>'[^']+'![A-Z]+\d+)=0,0,(?P=ref)\)$")
RE_IFBLANK = re.compile(r"^=IF\((?P<ref>'[^']+'![A-Z]+\d+)=\"\",\s*\"\",\s*(?P=ref)\)$")
RE_PLAIN = re.compile(r"^=(?P<ref>'\{Equipment Data Entry\}'![A-Z]+\d+)$")
RE_PCT = re.compile(r'^=IF\((?P<k>[A-Z]+\d+)="",\s*(?P<j>[A-Z]+\d+)/(?P<h>[A-Z]+\d+),\s*(?P<l>[A-Z]+\d+)/(?P=h)\)$')
RE_RATIO = re.compile(r'^=IF\((?P<l>[A-Z]+\d+)="","",(?P=l)/(?P<k>[A-Z]+\d+)\)$')
RE_SUB = re.compile(r'^=IF\((?P<a>[A-Z]+\d+)="",\s*"",\s*(?P<b>[A-Z]+\d+)-(?P<c>[A-Z]+\d+)\)$')


def hygiene(wb):
    counts = {}
    airflow_like = {"RTU Airflow", "MAU Airflow", "Fan Airflow", "VAV 1-20 Airflow", "Hoods"}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str) or not v.startswith("="):
                    continue
                new = None
                m = RE_IF0.match(v)
                if m and "RTU Airflow" not in v and "MAU Airflow" not in v and "Fan Airflow" not in v:
                    new = f'=IF({m["ref"]}="","",{m["ref"]})'; key = "blank shows \"\" not 0"
                elif RE_PLAIN.match(v):
                    ref = RE_PLAIN.match(v)["ref"]
                    new = f'=IF({ref}="","",{ref})'; key = "blank shows \"\" not 0"
                elif RE_PCT.match(v):
                    m = RE_PCT.match(v)
                    new = (f'=IF(OR({m["h"]}="",{m["h"]}=0),"",IF({m["k"]}="",IF({m["j"]}="","",'
                           f'{m["j"]}/{m["h"]}),{m["l"]}/{m["h"]}))'); key = "% guarded against blank/zero"
                elif RE_RATIO.match(v):
                    m = RE_RATIO.match(v)
                    new = f'=IF(OR({m["l"]}="",{m["k"]}="",{m["k"]}=0),"",{m["l"]}/{m["k"]})'; key = "ratio guarded"
                elif RE_SUB.match(v) and ws.title in ("RTU Data", "RTU Airflow"):
                    m = RE_SUB.match(v)
                    new = f'=IF({m["a"]}="","",{m["b"]}-N({m["c"]}))'; key = "subtraction tolerates blank"
                elif v.startswith("=IF(C") and "1-phase" in v and "746" in v and "COUNT(" not in v:
                    r = c.row
                    new = f'=IF(OR(COUNT(E{r+1}:G{r+1})=0,COUNT(E{r+2}:G{r+2})=0),"",{v[1:]})'; key = "BHP guarded"
                elif v.startswith("=IF(E") and "AVERAGE(E" in v and "*E" in v and ws.title.endswith("Data") or (v.startswith("=IF(E") and "AVERAGE(E" in v and "*E" in v and "Fan Data" in ws.title):
                    r = c.row
                    m2 = re.match(r'^=IF\(E(\d+)="",\s*"",\s*(B\d+)/\(?AVERAGE\(E\d+(?:,\s*F\d+,\s*G\d+|:G\d+)\)\)?\*(E\d+)\)$', v)
                    if m2:
                        new = f'=IF(OR(E{m2[1]}="",{m2[3]}="",{m2[2]}=""),"",{m2[2]}/AVERAGE(E{m2[1]}:G{m2[1]})*{m2[3]})'; key = "corrected FLA guarded"
                if new and new != v:
                    c.value = new
                    counts[(ws.title, key)] = counts.get((ws.title, key), 0) + 1
    for (t, k), n in sorted(counts.items()):
        log(f"{t}: {n} cells – {k}")
    # Building balance ratios
    bb = wb["Building Balance"]
    for r in range(7, 57):
        for col, e, d in (("G", "E", "C"), ("M", "K", "I")):
            if bb[f"{col}{r}"].value:
                bb[f"{col}{r}"].value = f'=IF(OR({e}{r}="",{d}{r}="",{d}{r}=0),"",{e}{r}/{d}{r})'
    log("Building Balance: % columns guarded against blank/zero design")


def harmonize_labels(wb):
    """Copy block-1 labels (and their style) into blocks where the cell is empty."""
    for name in ("RTU Data", "MAU Data", "Fan Data (EFs, TFs, etc.)"):
        ws = wb[name]; aa = anchors(ws); base = aa[0]; n = 0
        labels = {}
        for row in ws.iter_rows(min_row=base, max_row=base + 22, max_col=24):
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("="):
                    labels[(c.row - base, c.column)] = c
        for anc in aa[1:]:
            for (off, col), src in labels.items():
                dst = ws.cell(anc + off, col)
                if type(dst).__name__ == "MergedCell":
                    continue
                if dst.value is None:
                    dst.value = src.value; copy_style(src, dst); n += 1
        if n:
            log(f"{name}: filled {n} missing block labels from block 1")


def fix_misc(wb):
    ede = wb["{Equipment Data Entry}"]
    setf(ede, "P2", "=IF('{Project Information}'!E12=\"\", \"\", '{Project Information}'!E12)", "Technician header now reads Project Information E12")
    setf(ede, "P3", "=IF('{Project Information}'!E11=\"\", \"\", '{Project Information}'!E11)", "Date header now reads Project Information E11")
    toc = wb["ToC"]
    setf(toc, "C37", "Appendix B - Abbreviations", "typo fixed")
    for s in ("Summary - New", "Summary - (E)"):
        wb[s].print_area = "A1:M62"; log(f"{s}: print area extended to row 62 so all remark rows print")
    ab = wb["Abbreviations"]
    setf(ab, "H17", "kW = Kilowatt", "spelling")
    setf(ab, "B27", "kW = Kilowatt", "spelling")
    setf(ab, "H21", "V = Volts", "stray character removed")
    setf(ab, "B16", "d = Density (lbs/ft3)", "density units")
    setf(ab, "H39", "N/A = Not Applicable (state Not Available / Not Accessible where that is the reason)", "NEBB exception wording")
    setf(ab, "H45", "Not Avail. = Not Available", None)
    setf(ab, "H46", "Not Acc. = Not Accessible", None)
    ab["H45"].font = copy.copy(ab["H44"].font); ab["H46"].font = copy.copy(ab["H44"].font)
    ab["H45"].alignment = copy.copy(ab["H44"].alignment); ab["H46"].alignment = copy.copy(ab["H44"].alignment)
    for r in (45, 46):
        for c in range(8, 13):
            copy_style(ab.cell(44, c), ab.cell(r, c))
    log("Abbreviations: added NEBB exception notations (Not Available / Not Accessible)")
    # air-device type codes used by the new Type column
    for r, txt in ((45, "SD/RG/EG = Supply Diffuser / Return Grille / Exhaust Grille"), (46, "LSD/PSP = Linear Slot Diffuser / Perforated Supply Plenum")):
        ab[f"B{r}"].value = txt; copy_style(ab["B44"], ab[f"B{r}"])
    # leftover text under the certificate images
    for s in ("NEBB Cert ", "NEBB Frm Cert"):
        ws = wb[s]
        for coord in ("E13", "G13", "J13", "G16", "L25", "E31", "G31"):
            if ws[coord].value is not None:
                ws[coord].value = None
        log(f"{s}: removed leftover instrument text cells hidden under the certificate image")


def certification(wb):
    ws = wb["Certification"]
    txt = ws["C39"].value
    new = txt.replace("is a representation of system measurements", "is a record of system measurements")
    setf(ws, "C39", new, "certification statement now matches NEBB 5.2.2 verbatim")
    ws.merge_cells("C36:L37")
    ws["C36"].value = "NEBB Certified TAB Firm:  a2b accurate air balancing, llc  –  Firm Certification No. 3673"
    copy_style(ws["C34"], ws["C36"])
    # signature block
    thin = Side(style="thin")
    ws.merge_cells("C51:G56")
    ws["C50"].value = "NEBB Certified TAB Professional Stamp (signed and dated):"
    ws["C50"].font = Font(name=ws["C34"].font.name, size=10, bold=True)
    for r in range(51, 57):
        for c in range(3, 8):
            cell = ws.cell(r, c)
            cell.border = Border(left=thin if c == 3 else None, right=thin if c == 7 else None,
                                 top=thin if r == 51 else None, bottom=thin if r == 56 else None)
    ws["I52"].value = "Signature:"; ws["I52"].font = Font(name=ws["C34"].font.name, size=10, bold=True)
    ws.merge_cells("I53:L53"); ws["I53"].border = Border(bottom=thin)
    ws["I55"].value = "Date:"; ws["I55"].font = Font(name=ws["C34"].font.name, size=10, bold=True)
    ws.merge_cells("I56:L56"); ws["I56"].border = Border(bottom=thin)
    for r in range(50, 58):
        ws.row_dimensions[r].height = 13.35
    ws.print_area = "A1:M58"
    log("Certification: firm name/number line, stamp box, signature and date lines added (NEBB 5.2.2)")


def cover_report_date(wb):
    pi = wb["{Project Information}"]
    pi.merge_cells("B14:D14"); pi.merge_cells("E14:H14")
    pi["B14"].value = "Report Date"; copy_style(pi["B13"], pi["B14"]); copy_style(pi["E13"], pi["E14"])
    for c in range(3, 9):
        copy_style(pi.cell(13, c), pi.cell(14, c))
    pi["E14"].number_format = "m/d/yyyy"
    cv = wb["Cover Page"]
    cv.merge_cells("C26:F27"); cv.merge_cells("G26:L27")
    cv["C26"].value = "REPORT DATE"; cv["G26"].value = "=IF('{Project Information}'!E14=\"\",\"\",'{Project Information}'!E14)"
    for c in range(3, 13):
        for r in (26, 27):
            copy_style(cv.cell(24 if r == 26 else 25, c), cv.cell(r, c))
    cv["G26"].number_format = "mmmm d, yyyy"
    log("Cover Page / Project Information: Report Date added (NEBB 5.2.1)")


# --------------------------------------------------------------------------- #
# 3. Names and external links
# --------------------------------------------------------------------------- #
def cleanup_names(wb):
    removed = []
    for name in list(wb.defined_names):
        dn = wb.defined_names[name]
        if "#REF!" in dn.attr_text or re.search(r"\[\d+\]", dn.attr_text):
            del wb.defined_names[name]; removed.append(name)
    local = 0
    for ws in wb.worksheets:
        for name in list(ws.defined_names):
            dn = ws.defined_names[name]
            if "#REF!" in dn.attr_text or re.search(r"\[\d+\]", dn.attr_text):
                del ws.defined_names[name]; local += 1
    n_links = len(wb._external_links)
    wb._external_links = []
    log(f"Removed {len(removed)} broken/external workbook names, {local} sheet-scoped ones and {n_links} external link parts")


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #
STEPS = [fix_fan_data_totals, fix_mau_airflow_ak, fix_rtu_data_blocks, hygiene, harmonize_labels,
         fix_misc, certification, cover_report_date, cleanup_names]


def build(steps=None, out=OUT):
    if steps is None:
        from build_steps2 import STEPS2
        steps = STEPS + STEPS2
    wb = load_workbook(ORIG, keep_vba=True)
    for step in steps:
        print(f"[{step.__name__}]")
        step(wb)
    tmp = tempfile.mktemp(suffix=".xlsm")
    wb.save(tmp)
    lost = xlsm_parts.restore(ORIG, tmp, out, footer="&amp;C&amp;8Page &amp;P", skip_footer=("Cover Page",),
                              extra_sheet_sources={"ERV Data": "MAU Data", "ERV Airflow": "MAU Airflow",
                                                   "MAU Supply Methods": "MAU Airflow", "Narrative": "Summary - New"})
    os.remove(tmp)
    with open(os.path.join(ROOT, "docs", "build-log.txt"), "w") as fh:
        fh.write("\n".join(LOG) + "\n")
    print("written", out)
    return out


if __name__ == "__main__":
    build()
