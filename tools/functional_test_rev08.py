"""Functional test for the revision 08 TAB workbook (the static pressure profile tools/build_rev08.py rebuilds).

    python3 tools/functional_test_rev08.py "08 - a2b_Blank_TAB_Workbook <date>.xlsm"

Copies of the workbook are filled with openpyxl, recalculated with LibreOffice (soffice headless) and checked:

  0. blank template: no error cells; the unit-type table (labels per type)
  1. RTU full profile: Δ per component, fan inlet / discharge, TSP across a blow-through heat section, ESP, Unit ΔP,
     "Unit ESP actual" (L P+8) reads the ESP
  2. RTU 3-point (Capitola RTU-3 readings): entering, coil leaving = fan inlet, heat leaving = discharge
  3. MAU (Capitola MAU-9), EF, DOAS (6 components), DHU (desiccant after the coil)
  4. a notation (Not Acc.) at the fan inlet blanks TSP and Unit ΔP, not the ESP; a fan leaving static marked Not Acc.
     (blow-through) gives TSP from the discharge; no error cells anywhere
"""
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

SRC = sys.argv[1]
tmpdir = tempfile.mkdtemp()
ERRS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "Err:")
DASH = "—"
results = []


def recalc(path):
    outdir = os.path.join(tmpdir, "conv")
    os.makedirs(outdir, exist_ok=True)
    r = subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", outdir, path],
                       capture_output=True, text=True, timeout=900)
    conv = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".xlsx")
    if not os.path.exists(conv):
        raise SystemExit(f"soffice could not recalculate {path}: {r.stdout} {r.stderr}")
    shutil.move(conv, path)


def prepared(name, *fills):
    path = os.path.join(tmpdir, name + ".xlsm")
    shutil.copy(SRC, path)
    if fills:
        wb = load_workbook(path, keep_vba=True)
        for f in fills:
            f(wb)
        wb.save(path)
    recalc(path)
    return load_workbook(path, data_only=True)


def errors(wb):
    return [f"{ws.title}!{c.coordinate}={c.value}" for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith(ERRS)]


def check(section, name, got, exp, tol=1e-6):
    ok = (got == exp) or (isinstance(got, (int, float)) and isinstance(exp, (int, float))
                          and not isinstance(got, bool) and abs(got - exp) < tol)
    results.append((section, name, got, exp, ok))


def block(P):
    """Rows of the profile in the block at anchor P."""
    return dict(L=P + 19, R=P + 20, D=P + 21, U=P + 22, S=P + 23, T=P + 25)


SLOTS = "CDEFGH"
# (sheet, anchor, unit type, entering, readings of components 1-6)
CASES = {
    "rtu_full": ("RTUs", 4, "RTU", -0.5, [-0.6, -0.9, -1.0, None, 0.5, None]),
    "rtu_3pt": ("RTUs", 108, "RTU", -0.317, [None, -0.806, None, None, 0.514, None]),
    "doas": ("RTUs", 212, "DOAS", -0.3, [-0.4, -0.7, -1.1, -1.2, 0.9, 1.0]),  # Filter, Coil, Wheel, Reheat, Fan, Heat
    "dhu": ("RTUs", 316, "DHU", -1.26, [-1.53, -2.35, -2.6, None, 0.572, None]),
    "rtu_notation": ("RTUs", 420, "RTU", -0.5, [None, "Not Acc.", None, None, 0.5, None]),
    "rtu_fan_na": ("RTUs", 524, "RTU", -0.5, [-0.6, -0.9, None, "Not Acc.", 0.5, None]),
    "mau": ("MAUs", 4, "MAU", -0.355, [None, -0.921, 0.339, None, None, None]),
    "ef": ("Fans", 4, "EF", -0.2, [0.6, None, None, None, None, None]),
}


def fill(wb):
    for sheet, P, ut, ent, vals in CASES.values():
        ws, b = wb[sheet], block(P)
        ws[f"D{b['U']}"] = ut
        ws[f"B{b['R']}"] = ent
        for c, v in zip(SLOTS, vals):
            if v is not None:
                ws[f"{c}{b['R']}"] = v


# ----------------------------------------------------------------------------------------- 0: blank
blank = prepared("blank")
check("0 blank", "no error cells", errors(blank), [])
dd = blank["{Dropdowns}"]
table = [[dd[f"{c}{r}"].value for c in ["AH", "AI", "AJ", "AK", "AL", "AM", "AN", "AO"]] for r in range(2, 8)]
check("0 blank", "unit types", [t[0] for t in table], ["RTU", "DOAS", "DHU", "MAU", "ERV", "EF"])
check("0 blank", "RTU order", table[0][1:], ["RA / OA", "Filter", "Coil", "Reheat", "Fan", "Heat", DASH])
check("0 blank", "DHU order", table[2][1:], ["OA", "Filter", "Coil", "Desiccant", "Fan", "Heat", DASH])
rt = blank["RTUs"]
check("0 blank", "RTU block labels", [rt[f"{c}23"].value for c in "BCDEFGH"],
      ["RA / OA", "Filter", "Coil", "Reheat", "Fan", "Heat", DASH])
check("0 blank", "Fans block labels", [blank["Fans"][f"{c}23"].value for c in "BCD"], ["Inlet", "Fan", DASH])
check("0 blank", "blank results", [rt["E29"].value, rt["I29"].value, rt["M29"].value, rt["D27"].value], [None] * 4)

# ----------------------------------------------------------------------------------------- 1-4: filled
wb = prepared("filled", fill)
check("4 errors", "no error cells", errors(wb), [])


def res(case):
    sheet, P, *_ = CASES[case]
    ws, b = wb[sheet], block(P)
    return ws, b, {
        "delta": [ws[f"{c}{b['D']}"].value for c in SLOTS],
        "fanInlet": ws[f"D{b['S']}"].value, "discharge": ws[f"I{b['S']}"].value,
        "tsp": ws[f"E{b['T']}"].value, "esp": ws[f"I{b['T']}"].value, "unitDp": ws[f"M{b['T']}"].value,
    }


ws, b, r = res("rtu_full")
check("1 RTU full", "labels", [ws[f"{c}{b['L']}"].value for c in SLOTS], ["Filter", "Coil", "Reheat", "Fan", "Heat", DASH])
check("1 RTU full", "Δ", r["delta"], ["Δ -0.10", "Δ -0.30", "Δ -0.10", None, None, None])
check("1 RTU full", "fan inlet", r["fanInlet"], -1.0)
check("1 RTU full", "discharge", r["discharge"], 0.5)
check("1 RTU full", "TSP (heat after the fan)", r["tsp"], 1.5)
check("1 RTU full", "ESP", r["esp"], 1.0)
check("1 RTU full", "Unit ΔP", r["unitDp"], -0.5)
check("1 RTU full", "Unit ESP actual (L P+8)", ws["L12"].value, 1.0)

ws, b, r = res("rtu_3pt")
check("2 RTU 3-point", "Δ", r["delta"], [None] * 6)
check("2 RTU 3-point", "fan inlet", r["fanInlet"], -0.806)
check("2 RTU 3-point", "TSP", r["tsp"], 1.32)
check("2 RTU 3-point", "ESP", r["esp"], 0.831)
check("2 RTU 3-point", "Unit ΔP", r["unitDp"], -0.489)

ws, b, r = res("mau")
check("3 MAU", "labels", [ws[f"{c}{b['L']}"].value for c in SLOTS], ["Filter", "Burner", "Fan", DASH, DASH, DASH])
check("3 MAU", "Δ", r["delta"], [None, None, "Δ 1.26", None, None, None])
check("3 MAU", "TSP", r["tsp"], 1.26)
check("3 MAU", "ESP", r["esp"], 0.694)
check("3 MAU", "Unit ΔP", r["unitDp"], -0.566)

ws, b, r = res("ef")
check("3 EF", "TSP", r["tsp"], 0.8)
check("3 EF", "ESP", r["esp"], 0.8)
check("3 EF", "Unit ΔP (fan first)", r["unitDp"], 0.0)

ws, b, r = res("doas")
check("3 DOAS", "labels", [ws[f"{c}{b['L']}"].value for c in SLOTS], ["Filter", "Coil", "Wheel", "Reheat", "Fan", "Heat"])
check("3 DOAS", "Δ", r["delta"], ["Δ -0.10", "Δ -0.30", "Δ -0.40", "Δ -0.10", "Δ 2.10", "Δ 0.10"])
check("3 DOAS", "TSP (fan's own reading)", r["tsp"], 2.1)
check("3 DOAS", "ESP", r["esp"], 1.3)
check("3 DOAS", "Unit ΔP", r["unitDp"], -0.9)

ws, b, r = res("dhu")
check("3 DHU", "labels", [ws[f"{c}{b['L']}"].value for c in SLOTS], ["Filter", "Coil", "Desiccant", "Fan", "Heat", DASH])
check("3 DHU", "fan inlet (after the desiccant wheel)", r["fanInlet"], -2.6)
check("3 DHU", "TSP", r["tsp"], 0.572 + 2.6)
check("3 DHU", "ESP", r["esp"], 0.572 + 1.26)

ws, b, r = res("rtu_notation")
check("4 notation", "fan inlet shows the notation", r["fanInlet"], "Not Acc.")
check("4 notation", "TSP blank", r["tsp"], None)
check("4 notation", "Unit ΔP blank", r["unitDp"], None)
check("4 notation", "ESP", r["esp"], 1.0)

ws, b, r = res("rtu_fan_na")
check("4 notation", "fan leaving Not Acc.: TSP from the discharge", r["tsp"], 1.4)
check("4 notation", "fan leaving Not Acc.: heat Δ blank", r["delta"][4], None)

bad = [x for x in results if not x[4]]
for sec, name, got, exp, ok in results:
    print(f"{'PASS' if ok else 'FAIL'}  {sec}: {name}" + ("" if ok else f"  got {got!r}, expected {exp!r}"))
print(f"\n{len(results) - len(bad)} / {len(results)} passed")
sys.exit(1 if bad else 0)
