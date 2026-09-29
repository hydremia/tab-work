"""Hydronic TAB workbook, revision H01, built from the airside workbook (revision 05).

    python3 tools/build_hydronic.py     # newest 05 workbook -> H01 - a2b_Blank_Hydronic_Workbook <date>.xlsm

Requirements: docs/HYDRONIC_REQUIREMENTS.md (NEBB 9th Edition 5.3.13 - 5.3.24, Section 9).

Kept from the airside workbook, unchanged (the report-wide items of NEBB 5.2): Cover Page (title line
"HYDRONIC SYSTEMS" added), ToC (hydronic sections), Narrative, Summary - New / (E), {Project Information},
{Dropdowns} (hydronic lists added), Photos (hidden), Certification, NEBB Cert, NEBB Frm Cert,
Abbreviations (hydronic terms), Calibration.

Removed: every airside sheet (Equipment Summary, Building Balance, RTUs, MAUs, ERVs, Fans, Small Fans,
VAVs, Hoods, Traverses) and {Equipment Data Entry}.

Added (same look: header rows, label / value styles and page size of the airside unit pages):
  {Hydronic Data Entry}  pump design data (20 pumps), like {Equipment Data Entry}
  System Summary         per system: pump design / final GPM, sum of the valves, %, diversity, wide-open
                         valve count, VFD setpoint (NEBB 9.4 / 9.5.3)
  Pumps                  two pumps per page: unit, motor (as the fan pages), design / actual, pump test
                         (standing, shut-off, wide open, final: suction, discharge, head) - NEBB 5.3.13 / 5.3.14
  Valves                 one system per page, 38 valves: tag, serves, make / model, size, type, design,
                         initial and final GPM, setting, dP, %, wide open - NEBB 5.3.15 / 5.3.16
  Plant Equipment        chillers, towers, boilers, heat exchangers: one row per water circuit, design /
                         actual GPM and dP - NEBB 5.3.18 - 5.3.24
  Flow Measurements      ultrasonic readings - NEBB 5.3.17

As on the airside: the app writes inputs only, the workbook does the math (head from the gauges, %,
totals, the system summary). A valve's GPM is an input: it comes from the manufacturer's chart / Cv
(the app's valve library) or the valve tag, which a formula cannot look up.
"""
import copy
import datetime as dt
import glob
import os
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

sys.path.insert(0, os.path.dirname(__file__))
import xlsm_parts  # noqa: E402
from build_rev02 import copy_block, copy_style, merge  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STAMP = os.environ.get("TAB_BUILD_DATE", dt.date.today().strftime("%-m-%-d-%y"))
OUT = os.path.join(ROOT, f"H01 - a2b_Blank_Hydronic_Workbook {STAMP}.xlsm")
HDE = "'{Hydronic Data Entry}'"
DD = "'{Dropdowns}'"
PAGE = 52
N_PUMPS = 20
PUMP_BLOCK = 26          # two pumps per page
VALVE_PAGES = 25
VALVE_ROWS = 38          # per page (one system per page)
PLANT_PAGES = 2
PLANT_ROWS = 40
FLOW_ROWS = 24
N_SYSTEMS = 30
HDE_FIRST = 7            # first pump row on {Hydronic Data Entry}
AIRSIDE = ["Equipment Summary", "Building Balance", "RTUs", "MAUs", "ERVs", "Fans", "Small Fans", "VAVs", "Hoods",
           "Traverses", "{Equipment Data Entry}"]
LOG = []

LISTS = {  # name -> (column on {Dropdowns}, values)
    "Pump.Type": ("AU", ["Constant speed", "VFD", "Integrated variable speed"]),
    "Flow.Method": ("AV", ["Calibrated valve", "Flow meter", "Pump curve", "Equipment ΔP", "Ultrasonic", "Heat transfer"]),
    "DP.Units": ("AW", ["psi", "ft w.g."]),
    "Valve.Type": ("AX", ["F", "A", "S"]),
    "Balance.Method": ("AY", ["Proportional", "Stepwise", "Other"]),
    "Yes.No": ("AZ", ["Yes", "No", "N/A"]),
    "Plant.Type": ("BA", ["Chiller (water-cooled)", "Chiller (air-cooled)", "Cooling tower", "Hot-water boiler",
                          "Heat exchanger (water-water)", "Heat exchanger (steam-water)"]),
    "Circuit": ("BB", ["Evaporator", "Condenser", "Primary", "Secondary", "Water"]),
    "Wide.Open": ("BC", ["✓"]),
}


def log(m):
    LOG.append(m); print("  -", m)


class Look:
    """Cell styles taken from the airside Fan page, so the hydronic pages look the same."""

    def __init__(self, fans):
        self.fans = fans
        self.header = fans["B7"]      # blue section header, bold, centred, boxed
        self.label = fans["B8"]       # grey field label, bold, boxed
        self.value = fans["D8"]       # white value, left, boxed
        self.num = fans["L9"]         # white value, centred, number format 0
        self.col = fans["K8"]         # column caption (grey, bold)
        self.small = fans["B6"]       # small caption line under the system row
        self.remark_h = fans["B51"]
        self.remark = fans["B52"]

    def put(self, ws, rng, style, value=None, fmt=None, align=None, bold=None, size=None):
        cells = ws[rng] if ":" in rng else ((ws[rng],),)
        for row in cells:
            for c in row:
                if type(c).__name__ == "MergedCell":
                    continue
                copy_style(style, c)
                if fmt:
                    c.number_format = fmt
                if align or bold is not None or size:
                    c.alignment = Alignment(horizontal=align or c.alignment.horizontal, vertical="center",
                                            wrap_text=c.alignment.wrap_text)
                    f = copy.copy(c.font)
                    c.font = Font(name=f.name, size=size or f.sz, bold=f.b if bold is None else bold, color=f.color)
        if ":" in rng:
            merge(ws, rng)
        first = ws[rng.split(":")[0]]
        if value is not None:
            first.value = value
        return first


def guard(expr, *refs):
    """Text notations (N/A, Not Acc. ...) in an input read as blank, as in revision 05."""
    return f'IF(OR({",".join(f"ISTEXT({r})" for r in refs)}),"",{expr})'


def pct(num, den):
    return f'=IF(OR({den}="",{den}=0,{num}=""),"",{guard(f"{num}/{den}", num, den)})'


def link(ref):
    return f'=IF({ref}="","",{ref})'


def dv(ws, sqref, name, allow_text=True):
    d = DataValidation(type="list", formula1=name, allow_blank=True, showErrorMessage=not allow_text)
    d.sqref = sqref
    ws.add_data_validation(d)


def new_sheet(wb, name, fans, widths, scale=93):
    ws = wb.create_sheet(name)
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.sheet_format.defaultRowHeight = 13.35
    copy_block(fans, 1, 3, ws, 1)
    for m in fans.merged_cells.ranges:
        if m.max_row <= 3:
            merge(ws, str(m))
    ws.print_title_rows = "1:3"
    ws.page_margins = copy.copy(fans.page_margins)
    ws.page_setup.orientation = "portrait"; ws.page_setup.scale = scale
    ws.sheet_view.showGridLines = fans.sheet_view.showGridLines
    return ws


def widen_to_n(ws):
    """The airside pages end at column M; tables here use N too, so the header box, system row and remarks
    that end at M are stretched to N (merge and borders)."""
    for m in list(ws.merged_cells.ranges):
        if m.max_col == 13 and m.min_col < 13:
            rng = str(m)
            ws.unmerge_cells(rng)
            for r in range(m.min_row, m.max_row + 1):
                copy_style(ws.cell(r, 13), ws.cell(r, 14))
            merge(ws, f"{get_column_letter(m.min_col)}{m.min_row}:N{m.max_row}")


def page_breaks(ws, pages):
    for k in range(1, pages):
        ws.row_breaks.append(Break(id=3 + PAGE * k))


def remarks_block(look, ws, r, lines=2):
    """The fan page's Remarks header + lines (styles and merges)."""
    copy_block(look.fans, 51, 52, ws, r, values=False)
    for i in range(2, lines + 1):
        copy_block(look.fans, 52, 52, ws, r + i, values=False)
    ws[f"B{r}"] = "Remarks"


def rows_height(ws, r1, r2, h=13.35):
    for r in range(r1, r2 + 1):
        ws.row_dimensions[r].height = h


# --------------------------------------------------------------------------- #
def dropdowns(wb):
    dd = wb["{Dropdowns}"]
    for name, (col, values) in LISTS.items():
        dd[f"{col}1"] = name
        for i, v in enumerate(values):
            dd[f"{col}{2 + i}"] = v
        wb.defined_names[name] = DefinedName(name, attr_text=f"{DD}!${col}$2:${col}${1 + len(values)}")
    log("{Dropdowns}: hydronic lists " + ", ".join(f"{n} ({LISTS[n][0]})" for n in LISTS))


def data_entry(wb, look, air_ede):
    ws = wb.create_sheet("{Hydronic Data Entry}")
    widths = {"A": 3.43, "B": 10.71, "C": 12.14, "D": 12.14, "E": 12.43, "F": 13.86, "G": 13.86, "H": 9.0, "I": 9.0,
              "J": 10.0, "K": 7.57, "L": 9.43, "M": 9.0, "N": 8.43, "O": 6.86, "P": 13.57}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    copy_block(air_ede, 1, 5, ws, 1, c2=17)
    for m in air_ede.merged_cells.ranges:
        if m.max_row <= 5:
            merge(ws, str(m))
    ws["B5"] = "Design Data Pumps (NEBB 5.3.13 / 5.3.14)"
    ws["P5"] = None
    for col in ("P", "Q"):
        ws[f"{col}5"].style = "Normal"
    ws.unmerge_cells("B5:O5"); merge(ws, "B5:P5")
    copy_style(ws["O5"], ws["P5"])
    heads = ["Designation", "Service", "System", "Location", "Manufacturer", "Model / Size", "Design GPM",
             "Design Head (ft)", "Connected Load (GPM)", "Motor HP", "Pump / Motor RPM", "Impeller (in.)", "Voltage",
             "Phase", "Pump Type"]
    ws.row_dimensions[6].height = 40
    for i, h in enumerate(heads):
        c = ws.cell(6, 2 + i); copy_style(air_ede["B6"], c); c.value = h
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for k in range(N_PUMPS):
        r = HDE_FIRST + k
        ws.row_dimensions[r].height = 13.35
        for i in range(len(heads)):
            copy_style(air_ede.cell(7, 2 + min(i, 15)), ws.cell(r, 2 + i))
    ws[f"B{HDE_FIRST}"] = "P-1"
    last = HDE_FIRST + N_PUMPS - 1
    dv(ws, f"N{HDE_FIRST}:N{last}", "Motor.Voltage")
    dv(ws, f"O{HDE_FIRST}:O{last}", "Phase")
    dv(ws, f"P{HDE_FIRST}:P{last}", "Pump.Type")
    ws.print_area = f"A1:Q{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_margins = copy.copy(air_ede.page_margins)
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    # helper columns for the System Summary (outside the print area)
    ws["S6"] = "Final GPM (Pumps page)"; ws["T6"] = "Final head (Pumps page)"
    for k in range(N_PUMPS):
        P = pump_anchor(k)
        ws[f"S{HDE_FIRST + k}"] = link(f"Pumps!L{P + 5}")
        ws[f"T{HDE_FIRST + k}"] = link(f"Pumps!L{P + 6}")
    ws.column_dimensions["S"].hidden = True; ws.column_dimensions["T"].hidden = True
    log(f"{{Hydronic Data Entry}}: pump design data, rows {HDE_FIRST}-{last} ({N_PUMPS} pumps); "
        "hidden S/T link each pump's final GPM / head for the System Summary")
    return ws


def pump_anchor(k):
    return 4 + k * PUMP_BLOCK


def pumps(wb, look):
    fans = look.fans
    ws = new_sheet(wb, "Pumps", fans, {"A": 3.43, "B": 5.0, "C": 9.7, "D": 6.0, "E": 7.14, "F": 5.43, "G": 6.86, "H": 6.86, "I": 6.86,
                    "J": 6.86, "K": 6.86, "L": 6.86, "M": 6.86, "N": 3.14})
    for k in range(N_PUMPS):
        P = pump_anchor(k); e = HDE_FIRST + k
        E = lambda col: f"{HDE}!{col}{e}"
        rows_height(ws, P, P + PUMP_BLOCK - 1)
        copy_block(fans, 4, 5, ws, P, values=False)
        ws[f"B{P}"] = "Pump"; ws[f"D{P}"] = link(E("B")); ws[f"G{P}"] = "Service"; ws[f"I{P}"] = link(E("C"))
        # caption line
        r = P + 2
        look.put(ws, f"B{r}", look.small, "Pump type:")
        look.put(ws, f"D{r}:F{r}", look.small, link(E("P")), align="left")
        look.put(ws, f"G{r}", look.small, "System:")
        look.put(ws, f"H{r}:I{r}", look.small, link(E("D")), align="left")
        look.put(ws, f"J{r}", look.small, "Flow method:")
        look.put(ws, f"L{r}:M{r}", look.value, None, align="left", size=9)
        dv(ws, f"L{r}", "Flow.Method")
        # unit data / design and performance
        look.put(ws, f"B{P+3}:G{P+3}", look.header, "Unit Data")
        look.put(ws, f"I{P+3}:M{P+3}", look.header, "Design and Performance Data")
        unit = [("Manufacturer", link(E("F"))), ("Model / Size", link(E("G"))), ("Serial Number", None),
                ("Location", link(E("E"))), ("Service", f"=I{P}")]
        for i, (lab, val) in enumerate(unit):
            rr = P + 4 + i
            look.put(ws, f"B{rr}:C{rr}", look.label, lab)
            look.put(ws, f"D{rr}:G{rr}", look.value, val)
        look.put(ws, f"I{P+4}:J{P+4}", look.label, "Design Criteria")
        look.put(ws, f"K{P+4}", look.col, "Design", align="center")
        look.put(ws, f"L{P+4}", look.col, "Actual", align="center")
        look.put(ws, f"M{P+4}", look.col, "%", align="center")
        design = [("Flow (GPM)", E("H"), None, True), ("Head (ft w.g.)", E("I"), f"=M{P+16}", True),
                  ("Pump RPM", E("L"), None, False), ("Impeller (in.)", E("M"), None, False),
                  ("Connected Load (GPM)", E("J"), "—", False)]
        for i, (lab, d, a, withpct) in enumerate(design):
            rr = P + 5 + i
            look.put(ws, f"I{rr}:J{rr}", look.label, lab)
            look.put(ws, f"K{rr}", look.num, link(d), fmt="0.0" if "in." in lab else "0")
            look.put(ws, f"L{rr}", look.num, a, fmt="0.0" if "in." in lab or "ft" in lab else "0")
            look.put(ws, f"M{rr}", look.num, pct(f"L{rr}", f"K{rr}") if withpct else None, fmt="0%")
        # motor data: the fan page block, links repointed
        copy_block(fans, 14, 20, ws, P + 11, c1=2, c2=8)
        ws[f"D{P+13}"] = link(E("K"))
        ws[f"B{P+14}"] = link(E("N")); ws[f"C{P+14}"] = link(E("O"))
        look.put(ws, f"B{P+18}:D{P+18}", look.label, "VFD Hz / Speed Setting")
        look.put(ws, f"E{P+18}:G{P+18}", look.num, None, fmt="General")
        look.put(ws, f"B{P+19}:D{P+19}", look.label, "Final Setpoint(s)")
        look.put(ws, f"E{P+19}:G{P+19}", look.value, None)
        # pump test (psi at the gauges; head in ft w.g.)
        sg, dz = f"$L${P+17}", f"$L${P+18}"
        look.put(ws, f"I{P+11}:M{P+11}", look.header, "Pump Test (psi)")
        look.put(ws, f"I{P+12}:J{P+12}", look.label, "Condition")
        for col, cap in zip("KLM", ("Suction", "Discharge", "Head ft")):
            look.put(ws, f"{col}{P+12}", look.col, cap, align="center", size=9)
        for i, lab in enumerate(("Pump Off (standing)", "No Flow (shut-off)", "Wide Open", "Final")):
            rr = P + 13 + i
            look.put(ws, f"I{rr}:J{rr}", look.label, lab, size=9 if i < 2 else None)
            look.put(ws, f"K{rr}", look.num, None, fmt="0.0")
            if i == 0:
                look.put(ws, f"L{rr}", look.num, "—"); look.put(ws, f"M{rr}", look.num, None)
                continue
            look.put(ws, f"L{rr}", look.num, None, fmt="0.0")
            head = (f"(L{rr}-K{rr})*2.31/IF(AND(ISNUMBER({sg}),{sg}>0),{sg},1)+IF(ISNUMBER({dz}),{dz},0)")
            look.put(ws, f"M{rr}", look.num, f'=IF(OR(K{rr}="",L{rr}=""),"",{guard(head, f"K{rr}", f"L{rr}")})', fmt="0.0")
        look.put(ws, f"I{P+17}:K{P+17}", look.label, "Specific Gravity", size=9)
        look.put(ws, f"L{P+17}:M{P+17}", look.num, None, fmt="0.00")
        look.put(ws, f"I{P+18}:K{P+18}", look.label, "Gauge Elev. Δ (ft)", size=9)
        look.put(ws, f"L{P+18}:M{P+18}", look.num, None, fmt="0.0")
        look.put(ws, f"I{P+19}:K{P+19}", look.label, "Flow by Meter/Valve", size=9)
        look.put(ws, f"L{P+19}:M{P+19}", look.num, None, fmt="0")
        remarks_block(look, ws, P + 21, 2)
    last = pump_anchor(N_PUMPS) - 1
    page_breaks(ws, N_PUMPS // 2)
    ws.print_area = f"A1:N{last}"
    log(f"Pumps: {N_PUMPS} pumps, two per page ({PUMP_BLOCK} rows each, 'Pump' anchors in B from row 4); motor block "
        "and its corrected FLA / BHP formulas copied from the fan page; head = (discharge - suction) x 2.31 / SG "
        "+ gauge elevation difference")
    return ws


def table_header(look, ws, r, cols):
    """cols: list of (range_row1, text, second_row_text_or_None). A None second row merges both rows."""
    for rng, text, sub in cols:
        a, b = (rng.split(":") + [rng])[:2]
        c1, c2 = a, b
        if sub is None:
            look.put(ws, f"{c1}{r}:{c2}{r + 1}", look.header, text, size=9)
            ws[f"{c1}{r}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else:
            look.put(ws, f"{c1}{r}" if c1 == c2 else f"{c1}{r}:{c2}{r}", look.header, text, size=9)
            for col, t in sub:
                look.put(ws, f"{col}{r + 1}", look.col, t, align="center", size=9)


def valve_anchor(k):
    return 4 + k * PAGE


def valves(wb, look):
    ws = new_sheet(wb, "Valves", look.fans, {"A": 3.43, "B": 4.0, "C": 9.5, "D": 12.0, "E": 11.5, "F": 5.0, "G": 4.5,
                                             "H": 6.3, "I": 6.3, "J": 6.0, "K": 5.5, "L": 6.3, "M": 5.5, "N": 4.0},
                   scale=90)
    for k in range(VALVE_PAGES):
        P = valve_anchor(k)
        rows_height(ws, P, P + PAGE - 1)
        copy_block(look.fans, 4, 5, ws, P, values=False)
        ws[f"B{P}"] = "System"; ws[f"G{P}"] = "Service"
        r = P + 2
        look.put(ws, f"B{r}:C{r}", look.label, "Pump(s)", size=9)
        look.put(ws, f"D{r}:E{r}", look.value, None)
        look.put(ws, f"F{r}:G{r}", look.label, "Instrument", size=9)
        look.put(ws, f"H{r}:K{r}", look.value, None, size=9)
        look.put(ws, f"L{r}", look.label, "ΔP in", size=9)
        look.put(ws, f"M{r}:N{r}", look.value, None, align="center", size=9)
        dv(ws, f"M{r}", "DP.Units")
        r = P + 3
        look.put(ws, f"B{r}:C{r}", look.label, "Method", size=9)
        look.put(ws, f"D{r}:E{r}", look.value, None, size=9)
        dv(ws, f"D{r}", "Balance.Method")
        look.put(ws, f"F{r}:K{r}", look.label, "Memory stops set / valves marked", size=9)
        look.put(ws, f"L{r}:N{r}", look.value, None, align="center", size=9)
        dv(ws, f"L{r}", "Yes.No")
        h = P + 4
        table_header(look, ws, h, [("B", "No.", None), ("C", "Valve", None), ("D", "Serves", None),
                                   ("E", "Make / Model", None), ("F", "Size", None), ("G", "Type", None),
                                   ("H", "Design GPM", None), ("I", "Initial GPM", None),
                                   ("J:L", "Final", [("J", "Setting"), ("K", "ΔP"), ("L", "GPM")]),
                                   ("M", "%", None), ("N", "WO", None)])
        first = P + 6; lastrow = first + VALVE_ROWS - 1
        for rr in range(first, lastrow + 1):
            for col in "BCDE":
                look.put(ws, f"{col}{rr}", look.value, None, size=9)
            for col, fmt in zip("FGHIJKL", ("General", "General", "0.0", "0.0", "General", "0.00", "0.0")):
                look.put(ws, f"{col}{rr}", look.num, None, fmt=fmt, size=9)
            look.put(ws, f"M{rr}", look.num, f'=IF(OR(H{rr}="",H{rr}=0),"",IF(L{rr}="",IF(I{rr}="","",'
                                                   f'{guard(f"I{rr}/H{rr}", f"I{rr}", f"H{rr}")}),'
                                                   f'{guard(f"L{rr}/H{rr}", f"L{rr}", f"H{rr}")}))', fmt="0%", size=9)
            look.put(ws, f"N{rr}", look.num, None, size=9)
            ws[f"P{rr}"] = f'=IF($D${P}="","",$D${P})'          # system of the row, for the System Summary
        dv(ws, f"G{first}:G{lastrow}", "Valve.Type")
        dv(ws, f"N{first}:N{lastrow}", "Wide.Open")
        t = lastrow + 1
        look.put(ws, f"B{t}:G{t}", look.label, "Total", align="right")
        for col in "HIL":
            look.put(ws, f"{col}{t}", look.num, f'=IF(SUM({col}{first}:{col}{lastrow})=0,"",SUM({col}{first}:{col}{lastrow}))',
                     fmt="0.0", size=9)
        look.put(ws, f"J{t}:K{t}", look.label, None)
        look.put(ws, f"M{t}", look.num, f'=IF(OR(H{t}="",H{t}=0),"",IF(L{t}="",IF(I{t}="","",I{t}/H{t}),L{t}/H{t}))',
                 fmt="0%", size=9)
        look.put(ws, f"N{t}", look.num, f'=IF(COUNTA(N{first}:N{lastrow})=0,"",COUNTA(N{first}:N{lastrow}))', size=9)
        leg = ws[f"B{t+1}"]
        leg.value = ("Type: F = fixed orifice, A = adjustable orifice (flow from the manufacturer's chart at the setting "
                     "and ΔP), S = self-adjusting (flow from the valve tag, ΔP within its range). WO = wide open.")
        leg.font = Font(name=look.label.font.name, size=7, italic=True)
        merge(ws, f"B{t+1}:N{t+2}")
        leg.alignment = Alignment(wrap_text=True, vertical="top")
        remarks_block(look, ws, t + 3, 2)
    ws.column_dimensions["P"].hidden = True
    widen_to_n(ws)
    page_breaks(ws, VALVE_PAGES)
    ws.print_area = f"A1:N{valve_anchor(VALVE_PAGES) - 1}"
    log(f"Valves: {VALVE_PAGES} pages, one system per page ('System' anchors in B every {PAGE} rows from row 4), "
        f"{VALVE_ROWS} valves each (rows anchor+6 .. anchor+{5 + VALVE_ROWS}); % = final (else initial) / design; "
        "hidden P = the page's system for the System Summary")
    return ws


def simple_table_sheet(wb, look, name, title, widths, pages, rows, cols, header_rows, fmts, pct_cols, extra_line,
                       scale=93):
    ws = new_sheet(wb, name, look.fans, widths, scale=scale)
    for k in range(pages):
        P = 4 + k * PAGE
        rows_height(ws, P, P + PAGE - 1)
        look.put(ws, f"B{P}:N{P}", look.header, title)
        extra_line(ws, P + 1)
        table_header(look, ws, P + 2, header_rows)
        first = P + 4
        for rr in range(first, first + rows):
            for col in cols:
                if col in pct_cols:
                    num, den = pct_cols[col]
                    look.put(ws, f"{col}{rr}", look.num, pct(f"{num}{rr}", f"{den}{rr}"), fmt="0%", size=9)
                else:
                    look.put(ws, f"{col}{rr}", look.num if col in fmts else look.value, None,
                             fmt=fmts.get(col), size=9)
        remarks_block(look, ws, first + rows + 1, 2)
    widen_to_n(ws)
    page_breaks(ws, pages)
    ws.print_area = f"A1:N{3 + pages * PAGE}"
    return ws


def plant(wb, look):
    def extra(ws, r):
        look.put(ws, f"B{r}:C{r}", look.label, "ΔP in", size=9)
        look.put(ws, f"D{r}", look.value, None, align="center", size=9)
        dv(ws, f"D{r}", "DP.Units")
        look.put(ws, f"E{r}", look.label, "Instrument", size=9)
        look.put(ws, f"F{r}:N{r}", look.value, None, size=9)

    widths = {"A": 3.43, "B": 7.0, "C": 11.5, "D": 6.5, "E": 9.3, "F": 7.5, "G": 7.0, "H": 8.0, "I": 5.5, "J": 4.8,
              "K": 5.5, "L": 4.8, "M": 4.6, "N": 8.8}
    ws = simple_table_sheet(
        wb, look, "Plant Equipment", "Chillers, Cooling Towers, Boilers and Heat Exchangers (NEBB 5.3.18 – 5.3.24)",
        widths, PLANT_PAGES, PLANT_ROWS, "BCDEFGHIJKLMN",
        [("B", "Unit", None), ("C", "Type", None), ("D", "Service", None), ("E", "Mfr.", None),
         ("F", "Model", None), ("G", "Serial", None), ("H", "Circuit", None),
         ("I:J", "Design", [("I", "GPM"), ("J", "ΔP")]), ("K:L", "Actual", [("K", "GPM"), ("L", "ΔP")]),
         ("M", "%", None), ("N", "Method", None)],
        {"I": "0.0", "J": "0.0", "K": "0.0", "L": "0.0"}, {"M": ("K", "I")}, extra, scale=90)
    for k in range(PLANT_PAGES):
        f = 4 + k * PAGE + 4
        dv(ws, f"C{f}:C{f + PLANT_ROWS - 1}", "Plant.Type")
        dv(ws, f"H{f}:H{f + PLANT_ROWS - 1}", "Circuit")
        dv(ws, f"N{f}:N{f + PLANT_ROWS - 1}", "Flow.Method")
    log(f"Plant Equipment: {PLANT_PAGES} pages x {PLANT_ROWS} rows, one row per water circuit (a water-cooled chiller "
        "= evaporator + condenser rows); design / actual GPM and ΔP, %, flow method")
    return ws


def flow(wb, look):
    def extra(ws, r):
        look.put(ws, f"B{r}:C{r}", look.label, "Instrument", size=9)
        look.put(ws, f"D{r}:N{r}", look.value, None, size=9)

    widths = {"A": 3.43, "B": 4.0, "C": 11.0, "D": 11.0, "E": 6.0, "F": 8.0, "G": 6.0, "H": 8.0, "I": 7.0, "J": 6.5,
              "K": 6.3, "L": 6.3, "M": 5.5, "N": 7.0}
    ws = simple_table_sheet(
        wb, look, "Flow Measurements", "Ultrasonic Flow Measurement (NEBB 5.3.17)", widths, 1, FLOW_ROWS,
        "BCDEFGHIJKLMN",
        [("B", "No.", None), ("C", "System Served", None), ("D", "Location", None), ("E", "Pipe Size", None),
         ("F", "Pipe Material", None), ("G", "Wall (in.)", None), ("H", "Transducer", None), ("I", "Config.", None),
         ("J", "Spacing", None), ("K:L", "Flow GPM", [("K", "Design"), ("L", "Meas.")]), ("M", "%", None),
         ("N", "Notes", None)],
        {"B": "General", "E": "General", "G": "0.000", "K": "0.0", "L": "0.0"}, {"M": ("L", "K")}, extra)
    log(f"Flow Measurements: ultrasonic readings, {FLOW_ROWS} rows")
    return ws


def summary(wb, look):
    widths = {"A": 3.43, "B": 7.0, "C": 7.0, "D": 7.5, "E": 9.5, "F": 6.5, "G": 6.5, "H": 6.5, "I": 6.5, "J": 5.5,
              "K": 7.6, "L": 5.0, "M": 9.0, "N": 3.14}
    ws = new_sheet(wb, "System Summary", look.fans, widths)
    rows_height(ws, 4, 3 + PAGE)
    look.put(ws, "B4:M4", look.header, "Hydronic System Summary")
    table_header(look, ws, 5, [("B:C", "System", None), ("D", "Service", None), ("E", "Pump(s)", None),
                               ("F:G", "Pump GPM", [("F", "Design"), ("G", "Final")]),
                               ("H:I", "Valve GPM", [("H", "Design"), ("I", "Final")]),
                               ("J", "%", None), ("K", "Diversity", None), ("L", "WO", None),
                               ("M", "VFD Setpoint", None)])
    merge(ws, "B5:C6")
    VP = "Valves!$P$1:$P$1400"
    first = 7
    for i in range(N_SYSTEMS):
        r = first + i
        look.put(ws, f"B{r}:C{r}", look.value, None, size=9)
        look.put(ws, f"D{r}", look.value, None, size=9)
        look.put(ws, f"E{r}", look.value, None, size=9)
        sys_ = f"$B{r}"
        sumif = lambda rng, val: f'=IF({sys_}="","",IF(SUMIF({rng},{sys_},{val})=0,"",SUMIF({rng},{sys_},{val})))'
        look.put(ws, f"F{r}", look.num, sumif(f"{HDE}!$D${HDE_FIRST}:$D${HDE_FIRST + N_PUMPS - 1}",
                                               f"{HDE}!$H${HDE_FIRST}:$H${HDE_FIRST + N_PUMPS - 1}"), fmt="0", size=9)
        look.put(ws, f"G{r}", look.num, sumif(f"{HDE}!$D${HDE_FIRST}:$D${HDE_FIRST + N_PUMPS - 1}",
                                               f"{HDE}!$S${HDE_FIRST}:$S${HDE_FIRST + N_PUMPS - 1}"), fmt="0", size=9)
        look.put(ws, f"H{r}", look.num, sumif(VP, "Valves!$H$1:$H$1400"), fmt="0", size=9)
        look.put(ws, f"I{r}", look.num, sumif(VP, "Valves!$L$1:$L$1400"), fmt="0", size=9)
        look.put(ws, f"J{r}", look.num, pct(f"I{r}", f"H{r}"), fmt="0%", size=9)
        look.put(ws, f"K{r}", look.num, f'=IF(OR(F{r}="",H{r}=""),"",F{r}/H{r})', fmt="0.00", size=9)
        look.put(ws, f"L{r}", look.num, f'=IF({sys_}="","",COUNTIFS({VP},{sys_},Valves!$N$1:$N$1400,"<>"))', size=9)
        look.put(ws, f"M{r}", look.value, None, size=9)
    n = ws[f"B{first + N_SYSTEMS + 1}"]
    n.value = ("Pump GPM from {Hydronic Data Entry} / Pumps (by System); terminal GPM = sum of the Valves pages with "
               "the same System. Diversity = pump design GPM / terminal design GPM (NEBB 9.5.3). WO = valves recorded "
               "wide open: at least one wide-open path per system (NEBB 9.4). ±10 % of design (NEBB 9.4).")
    n.font = Font(name=look.label.font.name, size=7, italic=True); n.alignment = Alignment(wrap_text=True, vertical="top")
    merge(ws, f"B{first + N_SYSTEMS + 1}:M{first + N_SYSTEMS + 3}")
    ws.print_area = f"A1:N{3 + PAGE}"
    log(f"System Summary: {N_SYSTEMS} systems (names typed in B; they match {{Hydronic Data Entry}} System and the "
        "Valves page System): pump design / final, terminal design / final, %, diversity, wide-open count, VFD setpoint")
    return ws


# --------------------------------------------------------------------------- #
def cover(wb):
    """CERTIFIED / TEST, ADJUST, and BALANCE REPORT / HYDRONIC SYSTEMS: the title moves up into the free rows
    11-12 and the old title rows 13-14 carry the hydronic line (the photo box below stays where it is)."""
    ws = wb["Cover Page"]
    title = ws["B13"]
    for r in (11, 12):
        ws.row_dimensions[r].height = 13.35
        for col in range(2, 14):
            copy_style(ws.cell(r + 2, col), ws.cell(r, col))
    merge(ws, "B11:M12")
    ws["B11"] = title.value
    ws["B13"] = "HYDRONIC SYSTEMS"
    log("Cover Page: 'TEST, ADJUST, and BALANCE REPORT' moved to rows 11-12, 'HYDRONIC SYSTEMS' in rows 13-14")


def toc(wb):
    ws = wb["ToC"]
    entries = ["Table of Contents", "Report Summary - Narrative and Remarks", "Hydronic System Summary", "Pumps",
               "Balancing Valves", "Plant Equipment and Flow Measurements", "Appendix A - Certifications",
               "Appendix B - Abbreviations", "Appendix C - Instrument Calibration", "Piping Schematic(s)",
               "Site Photos"]
    for i in range(14):
        r = 10 + 3 * i
        if i < len(entries):
            ws[f"C{r}"] = entries[i]; ws[f"K{r}"] = f"page {i + 2}"
        else:
            ws[f"C{r}"] = None; ws[f"K{r}"] = None
    log("ToC: " + "; ".join(entries) + " (page numbers are written by the SyncToCPageCounts macro, TABHydronic.bas)")


ABBR_B = ["BHP = Brake Horsepower", "Btu = British Thermal Unit", "CBV = Calibrated Balancing Valve",
          "Cv = Valve Flow Coefficient (GPM at 1 psi ΔP)", "∆ = Difference (final – initial)",
          "∆P = Differential Pressure", "EWT / LWT = Entering / Leaving Water Temp.",
          "Ft = Foot", "ft w.g. = Feet of Water Gauge (1 psi = 2.31 ft w.g.)", "FR = Frame",
          "GPM = Gallons per Minute", "HW / CHW / CW = Hot / Chilled / Cond. Water",
          "HX = Heat Exchanger", "Hr = Hour", "Imp. = Impeller Diameter", "kW = Kilowatt", "OD = Outer Diameter",
          "PICV = Pressure-Independent Control Valve", "psi = Pounds per Square Inch", "psig = psi, Gauge",
          "RPM = Revolutions per Minute", "SF = Service Factor", "SFA = Service Factor Amperage",
          "SG = Specific Gravity (water = 1.00)", "T = Temperature (°F)", "TDH = Total Dynamic Head (ft w.g.)",
          "VFD = Variable Frequency Drive", "VSD = Variable Speed Drive", "WO = Wide Open",
          "F / A / S = Fixed / Adjustable / Self-Adj. Valve", "Shut-off Head = Head at No Flow",
          "Standing Pressure = System Pressure, Pumps Off", "Diversity = Pump GPM / Sum of Terminal GPM", "", "", ""]
ABBR_TEMP = ["Temperature Nomenclatures", "°C = Degrees in Celsius", "°F = Degrees in Fahrenheit",
             "EWT = Entering Water Temperature", "LWT = Leaving Water Temperature", "∆T = Temperature Difference",
             "EAT / LAT = Entering / Leaving Air Temp.", "", "", "", ""]


def abbreviations(wb):
    ws = wb["Abbreviations"]
    ws["B10"] = "Hydronic and Motor Nomenclatures"
    for i, t in enumerate(ABBR_B):
        ws[f"B{11 + i}"] = t or None
    for i, t in enumerate(ABBR_TEMP):
        ws[f"H{25 + i}"] = t or None
    ws["H43"] = "CBV = Calibrated Balancing Valve"
    log("Abbreviations: hydronic list in column B, temperature list in H25:H31 (electrical and reporting kept)")


def order(wb):
    names = ["Cover Page", "ToC", "Narrative", "Summary - New", "Summary - (E)", "{Project Information}",
             "{Hydronic Data Entry}", "{Dropdowns}", "System Summary", "Pumps", "Valves", "Plant Equipment",
             "Flow Measurements", "Photos", "Certification", "NEBB Cert ", "NEBB Frm Cert", "Abbreviations", "Calibration"]
    wb._sheets = [wb[n] for n in names]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Cover Page"
    log("sheet order: " + ", ".join(names))


def check_refs(wb):
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and any(f"'{s}'!" in v or f"{s}!" in v for s in AIRSIDE):
                    bad.append(f"{ws.title}!{c.coordinate}")
    if bad:
        raise SystemExit("formulas still point at removed airside sheets: " + ", ".join(bad[:20]))
    log("no formula refers to a removed airside sheet")


def build(src=None, out=OUT):
    src = src or sorted(glob.glob(os.path.join(ROOT, "05 - a2b_Blank_TAB_Workbook *.xlsm")))[-1]
    print("source:", os.path.basename(src))
    LOG.append(f"Source: {os.path.basename(src)}")
    wb = load_workbook(src, keep_vba=True)
    look = Look(wb["Fans"])
    air_ede = wb["{Equipment Data Entry}"]
    dropdowns(wb)
    data_entry(wb, look, air_ede)
    summary(wb, look)
    pumps(wb, look)
    valves(wb, look)
    plant(wb, look)
    flow(wb, look)
    cover(wb)
    toc(wb)
    abbreviations(wb)
    for name in AIRSIDE:
        wb.remove(wb[name])
    log("removed: " + ", ".join(AIRSIDE))
    order(wb)
    check_refs(wb)
    tmp = tempfile.mktemp(suffix=".xlsm")
    wb.save(tmp)
    titles = {"Pumps": "Pump Report", "Valves": "Balancing Valve Report", "Plant Equipment": "Plant Equipment Report",
              "Flow Measurements": "Flow Measurement Report", "System Summary": "Hydronic System Summary"}
    xlsm_parts.restore(src, tmp, out, footer="&amp;C&amp;8Page &amp;P", skip_footer=("Cover Page",),
                       extra_sheet_sources={n: "Fans" for n in titles} | {"{Hydronic Data Entry}": "{Equipment Data Entry}"},
                       header_overrides=titles)
    os.remove(tmp)
    with open(os.path.join(ROOT, "docs", "build-log-hydronic-h01.txt"), "w") as fh:
        fh.write("\n".join(LOG) + "\n")
    print("written", out)
    return out


if __name__ == "__main__":
    build()
