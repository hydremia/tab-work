"""Data-sheet to Airflow/Data-Entry link verification.  python3 tools/verify_links.py <workbook.xlsm>"""
import sys
from openpyxl import load_workbook
import re
wb = load_workbook(sys.argv[1], keep_vba=True)
def anchors(ws, label='System'): return [c.row for c in ws['B'] if c.value==label]
def refs(f, sheet):
    return [int(m.group(1)) for m in re.finditer(re.escape(sheet)+r"'!\$?[A-Z]+\$?(\d+)", f or "")]
def ede_refs(f): return refs(f, "{Equipment Data Entry}")

def check_pair(dname, aname, total_off, oa_off=None, ede_base=None):
    d=wb[dname]; a=wb[aname]; da=anchors(d); aa=anchors(a)
    print(f"\n## {dname} ({len(da)} blocks) -> {aname} ({len(aa)} blocks)")
    for n,anc in enumerate(da):
        # find the Total Airflow row in this block: label 'Total Airflow' in col I
        prob=[]
        for r in range(anc, anc+23):
            if d.cell(r,9).value=='Total Airflow':
                for col in (11,12):
                    f=d.cell(r,col).value; got=refs(f, aname)
                    exp=aa[n]+total_off if n<len(aa) else None
                    if got!=[exp] and got!=[exp,exp]: prob.append(f"{d.cell(r,col).coordinate} Total refs {got} expected {exp}: {str(f)[:70]}")
            if oa_off and d.cell(r,9).value=='Outside Airflow':
                for col in (11,12):
                    f=d.cell(r,col).value; got=refs(f, aname)
                    exp=aa[n]+oa_off if n<len(aa) else None
                    if got!=[exp] and got!=[exp,exp]: prob.append(f"{d.cell(r,col).coordinate} OA refs {got} expected {exp}: {str(f)[:70]}")
        # EDE row check: all EDE refs within block must be ede_base+n
        if ede_base:
            for row in d.iter_rows(min_row=anc, max_row=anc+22):
                for c in row:
                    if isinstance(c.value,str) and '{Equipment Data Entry}' in c.value:
                        got=set(ede_refs(c.value))
                        if got!={ede_base+n}: prob.append(f"{c.coordinate} EDE refs {got} expected {ede_base+n}: {c.value[:60]}")
        # system name cell D
        if prob: print(f" block {n+1} @row {anc}:"); [print("   ",p) for p in prob]

check_pair('RTU Data','RTU Airflow', total_off=33, oa_off=44, ede_base=7)
check_pair('MAU Data','MAU Airflow', total_off=19, ede_base=52)
check_pair('Fan Data (EFs, TFs, etc.)','Fan Airflow', total_off=19, ede_base=81)

# Airflow sheets: EDE refs per block
for aname,base in [('RTU Airflow',7),('MAU Airflow',52),('Fan Airflow',81),('Hoods',126)]:
    a=wb[aname]; aa=anchors(a); print(f"\n## {aname} EDE refs")
    for n,anc in enumerate(aa):
        end=aa[n+1] if n+1<len(aa) else a.max_row
        for row in a.iter_rows(min_row=anc, max_row=end-1):
            for c in row:
                if isinstance(c.value,str) and '{Equipment Data Entry}' in c.value:
                    got=set(ede_refs(c.value))
                    if got!={base+n}: print(f"   block {n+1} {c.coordinate}: {got} expected {base+n}")
# Building Balance
bb=wb['Building Balance']; rd=wb['RTU Data']; md=wb['MAU Data']; fd=wb['Fan Data (EFs, TFs, etc.)']
rda=anchors(rd); mda=anchors(md); fda=anchors(fd)
print("\n## Building Balance")
for r in range(7,57):
    b=bb.cell(r,2).value; c=bb.cell(r,3).value; h=bb.cell(r,8).value; i=bb.cell(r,9).value
    print(f" row {r}: B={b[:55] if b else b} | C={c[:45] if c else c} | H={h[:55] if h else h} | I={i[:50] if i else i}")
