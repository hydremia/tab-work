"""Functional test for the hydronic TAB workbook (revision H01).

    python3 tools/functional_test_hydronic.py ["H01 - a2b_Blank_Hydronic_Workbook <date>.xlsm"]

Recalculates copies of the workbook with LibreOffice (soffice headless) and checks:
  0. blank template: no error cells, no formula refers to a removed airside sheet
  A. sample data: pump head from the gauges (SG and gauge elevation), design / actual / %, the fan-page motor
     formulas (corrected FLA, BHP), valve %, totals and wide-open count, the System Summary (pump and valve
     sums by system, %, diversity), plant and flow-measurement %
  B. N/A notations in numeric inputs: dependent cells blank, totals skip them, no error cells
"""
import glob
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = sys.argv[1] if len(sys.argv) > 1 else sorted(glob.glob(os.path.join(ROOT, "H01 - a2b_Blank_Hydronic_Workbook *.xlsm")))[-1]
ERRS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "Err:")
tmpdir = tempfile.mkdtemp()
results = []


def recalc(path):
    outdir = os.path.join(tmpdir, "conv")
    os.makedirs(outdir, exist_ok=True)
    r = subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", outdir, path],
                       capture_output=True, text=True, timeout=900)
    conv = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".xlsx")
    if not os.path.exists(conv):
        raise SystemExit(f"soffice could not recalculate {path}: {r.stdout} {r.stderr}")
    return load_workbook(conv, data_only=True)


def prepared(name, fill=None):
    path = os.path.join(tmpdir, name + ".xlsm")
    shutil.copy(SRC, path)
    if fill:
        wb = load_workbook(path, keep_vba=True)
        fill(wb)
        wb.save(path)
    return recalc(path)


def check(section, name, got, expected, tol=0.01):
    if isinstance(expected, float) and isinstance(got, (int, float)):
        ok = abs(got - expected) <= tol
    else:
        ok = got == expected
    results.append((section, name, got, expected, ok))


def errors(wb):
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERRS):
                    out.append(f"{ws.title}!{c.coordinate}={c.value}")
    return out


P1, P2 = 4, 30          # first two pump blocks
V1, V2 = 4, 56          # first two valve pages
HDE = "{Hydronic Data Entry}"


def sample(wb, notation=False):
    e = wb[HDE]
    e["B7"], e["C7"], e["D7"], e["E7"], e["F7"], e["G7"] = "P-1", "Chilled Water", "CHW", "Mech Rm", "B&G", "e-1510 3BC"
    e["H7"], e["I7"], e["J7"], e["K7"], e["L7"], e["M7"], e["N7"], e["O7"], e["P7"] = 200, 60, 180, 7.5, 1750, 7.5, 460, "3-phase", "VFD"
    e["B8"], e["C8"], e["D8"], e["H8"], e["I8"] = "P-2", "Hot Water", "HW", 120, 40
    p = wb["Pumps"]
    p[f"L{P1 + 5}"] = 190                                     # final flow from the curve
    p[f"K{P1 + 13}"] = 12                                     # standing
    p[f"K{P1 + 14}"], p[f"L{P1 + 14}"] = 10, 40               # shut-off: 30 psi -> 69.3 ft
    p[f"K{P1 + 15}"], p[f"L{P1 + 15}"] = 8, 30                # wide open: 22 psi -> 50.82 ft
    p[f"K{P1 + 16}"], p[f"L{P1 + 16}"] = ("N/A" if notation else 9), 35   # final: 26 psi -> 60.06 ft
    p[f"E{P1 + 14}"] = 9.6                                    # FLA
    p[f"E{P1 + 16}"], p[f"F{P1 + 16}"], p[f"G{P1 + 16}"] = 462, 460, 458      # measured volts
    p[f"E{P1 + 17}"], p[f"F{P1 + 17}"], p[f"G{P1 + 17}"] = 8.0, "Not Acc." if notation else 8.2, 8.1  # amps
    p[f"L{P2 + 5}"] = 130
    p[f"K{P2 + 16}"], p[f"L{P2 + 16}"] = 15, 30               # 15 psi
    p[f"L{P2 + 17}"] = 1.10                                   # SG (glycol): 15*2.31/1.1 = 31.5
    p[f"L{P2 + 18}"] = 1.5                                    # gauge elevation: +1.5 ft -> 33.0
    v = wb["Valves"]
    v[f"D{V1}"], v[f"I{V1}"] = "CHW", "Chilled Water"
    rows = [(10, 9, 11), (20, 17, 19.5), (30, 36, "N/A" if notation else 31)]
    for i, (d, ini, fin) in enumerate(rows):
        r = V1 + 6 + i
        v[f"B{r}"], v[f"C{r}"], v[f"H{r}"], v[f"I{r}"], v[f"L{r}"] = i + 1, f"CBV-{i + 1}", d, ini, fin
    v[f"N{V1 + 6}"] = "✓"
    v[f"D{V2}"] = "HW"
    v[f"H{V2 + 6}"], v[f"L{V2 + 6}"] = 50, 55
    v[f"H{V2 + 7}"], v[f"L{V2 + 7}"], v[f"N{V2 + 7}"] = 70, 66, "✓"
    s = wb["System Summary"]
    s["B7"], s["B8"] = "CHW", "HW"
    pl = wb["Plant Equipment"]
    pl["B8"], pl["C8"], pl["H8"], pl["I8"], pl["K8"] = "CH-1", "Chiller (water-cooled)", "Evaporator", 240, 228
    fl = wb["Flow Measurements"]
    fl["K8"], fl["L8"] = 200, 210


def main():
    print("workbook:", os.path.basename(SRC))
    blank = prepared("blank")
    errs = errors(blank)
    check("0", "blank template: error cells", len(errs), 0)
    names = [ws.title for ws in blank.worksheets]
    check("0", "sheets", [n for n in names if n in ("System Summary", "Pumps", "Valves", "Plant Equipment",
                                                     "Flow Measurements", HDE)],
          [HDE, "System Summary", "Pumps", "Valves", "Plant Equipment", "Flow Measurements"])
    check("0", "no airside sheets", [n for n in names if n in ("RTUs", "Fans", "VAVs", "{Equipment Data Entry}")], [])

    wb = prepared("sample", sample)
    p, v, s = wb["Pumps"], wb["Valves"], wb["System Summary"]
    check("A", "pump 1 designation link", p[f"D{P1}"].value, "P-1")
    check("A", "pump 1 service link", p[f"I{P1}"].value, "Chilled Water")
    check("A", "pump 1 type / system caption", (p[f"D{P1 + 2}"].value, p[f"H{P1 + 2}"].value), ("VFD", "CHW"))
    check("A", "shut-off head (30 psi)", p[f"M{P1 + 14}"].value, 69.3)
    check("A", "wide-open head (22 psi)", p[f"M{P1 + 15}"].value, 50.82)
    check("A", "final head (26 psi)", p[f"M{P1 + 16}"].value, 60.06)
    check("A", "actual head = final head", p[f"L{P1 + 6}"].value, 60.06)
    check("A", "design head link", p[f"K{P1 + 6}"].value, 60)
    check("A", "head %", p[f"M{P1 + 6}"].value, 1.001)
    check("A", "flow design / actual / %", (p[f"K{P1 + 5}"].value, p[f"L{P1 + 5}"].value), (200, 190))
    check("A", "flow %", p[f"M{P1 + 5}"].value, 0.95)
    check("A", "motor HP link", p[f"D{P1 + 13}"].value, 7.5)
    check("A", "volts / phase links", (p[f"B{P1 + 14}"].value, p[f"C{P1 + 14}"].value), (460, "3-phase"))
    check("A", "corrected FLA = FLA x 460 / avg V", p[f"D{P1 + 15}"].value, 9.6 * 460 / 460)
    check("A", "estimated BHP (3-phase)", p[f"G{P1 + 15}"].value, 460 * 8.1 * 0.8 * 0.9 * 1.732 / 746, tol=0.02)
    check("A", "pump 2: SG 1.10 and +1.5 ft", p[f"M{P2 + 16}"].value, 15 * 2.31 / 1.1 + 1.5)
    check("A", "valve % (final / design)", v[f"M{V1 + 6}"].value, 1.1)
    check("A", "valve total design / initial / final", (v[f"H{V1 + 44}"].value, v[f"I{V1 + 44}"].value, v[f"L{V1 + 44}"].value),
          (60, 62, 61.5))
    check("A", "valve total %", v[f"M{V1 + 44}"].value, 61.5 / 60)
    check("A", "wide-open count", v[f"N{V1 + 44}"].value, 1)
    check("A", "summary CHW pump design / final", (s["F7"].value, s["G7"].value), (200, 190))
    check("A", "summary CHW valve design / final", (s["H7"].value, s["I7"].value), (60, 61.5))
    check("A", "summary CHW %", s["J7"].value, 61.5 / 60)
    check("A", "summary CHW diversity", s["K7"].value, 200 / 60)
    check("A", "summary CHW wide open", s["L7"].value, 1)
    check("A", "summary HW pump / valves", (s["F8"].value, s["G8"].value, s["H8"].value, s["I8"].value), (120, 130, 120, 121))
    check("A", "summary empty row blank", (s["F9"].value, s["J9"].value, s["L9"].value), (None, None, None))
    check("A", "plant %", wb["Plant Equipment"]["M8"].value, 0.95)
    check("A", "flow measurement %", wb["Flow Measurements"]["M8"].value, 1.05)
    check("A", "error cells", errors(wb), [])

    wb = prepared("notation", lambda w: sample(w, notation=True))
    p, v = wb["Pumps"], wb["Valves"]
    check("B", "final head blank when suction is N/A", p[f"M{P1 + 16}"].value, None)
    check("B", "N/A input still shows", p[f"K{P1 + 16}"].value, "N/A")
    check("B", "BHP from the two readable amp legs (8.0, 8.1; one Not Acc.)", p[f"G{P1 + 15}"].value,
          460 * 8.05 * 0.8 * 0.9 * 1.732 / 746, tol=0.02)
    check("B", "valve % blank when the final reading is N/A (as airside)", v[f"M{V1 + 8}"].value, None)
    check("B", "valve final total skips N/A", v[f"L{V1 + 44}"].value, 30.5)
    check("B", "error cells", errors(wb), [])

    fails = [r for r in results if not r[4]]
    for sec, name, got, exp, ok in results:
        print(f"{'PASS' if ok else 'FAIL'}  {sec}  {name}" + ("" if ok else f": got {got!r}, expected {exp!r}"))
    print(f"\n{len(results) - len(fails)} PASS, {len(fails)} FAIL")
    shutil.rmtree(tmpdir, ignore_errors=True)
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
