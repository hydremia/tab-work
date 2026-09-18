"""Revision 03: equipment summary page, small exhaust fans (two per page), duct-traverse point grids.

    python3 tools/build_rev03.py     # reads the newest 02 workbook, writes 03 - a2b_Blank_TAB_Workbook <date>.xlsm
"""
import copy
import datetime as dt
import glob
import os
import re
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, os.path.dirname(__file__))
import xlsm_parts  # noqa: E402
from build_rev02 import (PAGE, EDE, DD, THIN, copy_block, copy_style, merge, merged_set, outlet_rows, total_row,  # noqa: E402
                         remarks, dv, new_sheet, anchors)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = sorted(glob.glob(os.path.join(ROOT, "02 - a2b_Blank_TAB_Workbook *.xlsm")))[-1]
STAMP = os.environ.get("TAB_BUILD_DATE", dt.date.today().strftime("%-m-%-d-%y"))
OUT = os.path.join(ROOT, f"03 - a2b_Blank_TAB_Workbook {STAMP}.xlsm")
EDE_SMALL = 233          # Data Entry rows for small exhaust fans
N_SMALL = 40
SF_BLOCK = 24            # rows per small fan (two per page)
TRAV_BLOCK = 15          # rows per traverse (three per page)
TRAV_PAGES = 16
LOG = []


def log(m):
    LOG.append(m); print("  -", m)


def ede_link(col, row):
    return f'=IF({EDE}!{col}{row}="","",{EDE}!{col}{row})'


# --------------------------------------------------------------------------- #
# 1. Data Entry: small exhaust fan section
# --------------------------------------------------------------------------- #
def data_entry_small_fans(wb):
    ws = wb["{Equipment Data Entry}"]
    top = EDE_SMALL - 2
    merge(ws, f"B{top}:O{top}")
    ws[f"B{top}"] = "Design Data Small Exhaust Fans (direct drive, under 1/6 hp – NEBB 5.3.6)"
    copy_style(ws["B124"], ws[f"B{top}"])
    heads = ["Designation", "Area Served", "Location", "Manufacturer", "Model Number", "HP", "Voltage", "Phase", "Design CFM"]
    for i, h in enumerate(heads):
        c = ws.cell(top + 1, 2 + i); c.value = h; copy_style(ws.cell(125, 2 + min(i, 6)), c)
    ws.row_dimensions[top + 1].height = ws.row_dimensions[125].height
    for r in range(EDE_SMALL, EDE_SMALL + N_SMALL):
        for i in range(len(heads)):
            copy_style(ws.cell(126, 2 + min(i, 6)), ws.cell(r, 2 + i))
        ws.row_dimensions[r].height = ws.row_dimensions[126].height
    ws[f"B{EDE_SMALL}"] = "EF-S1"
    ws.print_area = f"A1:S{EDE_SMALL + N_SMALL}"
    log("{Equipment Data Entry}: small exhaust fan section (40 rows) added")


# --------------------------------------------------------------------------- #
# 2. Small Fans sheet – two per page
# --------------------------------------------------------------------------- #
def small_fans(wb):
    fans = wb["Fans"]; rtu = wb["RTUs"]
    ws = new_sheet(wb, "Small Fans", fans, fans, fans)
    ws.sheet_format.defaultColWidth = 6.86
    instr, phase = [], []
    for u in range(N_SMALL):
        P = 4 + u * SF_BLOCK; e = EDE_SMALL + u
        copy_block(fans, 4, 5, ws, P, values=False)
        ws[f"B{P}"] = "System"; ws[f"D{P}"] = ede_link("B", e); ws[f"G{P}"] = "Service"; ws[f"I{P}"] = ede_link("C", e)
        copy_block(fans, 7, 12, ws, P + 2, values=False)      # Unit Data header + 5 rows, right side performance rows
        ws[f"B{P+2}"] = "Unit Data"; ws[f"I{P+2}"] = "Design and Performance Data"
        ws[f"B{P+3}"] = "Manufacturer"; ws[f"D{P+3}"] = ede_link("E", e); ws[f"I{P+3}"] = "Design Criteria"; ws[f"K{P+3}"] = "Design"; ws[f"L{P+3}"] = "Actual"
        ws[f"B{P+4}"] = "Model Number"; ws[f"D{P+4}"] = ede_link("F", e); ws[f"I{P+4}"] = "Total Airflow"
        ws[f"B{P+5}"] = "Serial Number"; ws[f"I{P+5}"] = "Unit ESP"
        ws[f"B{P+6}"] = "Area Served"; ws[f"D{P+6}"] = ede_link("C", e); ws[f"I{P+6}"] = "Fan RPM"
        ws[f"B{P+7}"] = "Location"; ws[f"D{P+7}"] = ede_link("D", e); ws[f"I{P+7}"] = "Speed Setting"
        for r in (P + 5, P + 6, P + 7):
            for c in range(9, 14):
                copy_style(fans.cell(10, c), ws.cell(r, c))
            merge(ws, f"I{r}:J{r}")
        # motor line and instrument / final settings line
        for r, (lab, right) in ((P + 8, ("HP / V / Ph", "Measured Amps")), (P + 9, ("Instrument", "Final Settings"))):
            for c in range(2, 8):
                copy_style(fans.cell(8, c), ws.cell(r, c))
            for c in range(9, 14):
                copy_style(fans.cell(10, c), ws.cell(r, c))
            merge(ws, f"B{r}:C{r}"); merge(ws, f"I{r}:J{r}"); merge(ws, f"K{r}:M{r}")
            ws[f"B{r}"] = lab; ws[f"I{r}"] = right
        ws[f"D{P+8}"] = ede_link("G", e); ws[f"E{P+8}"] = ede_link("H", e); ws[f"F{P+8}"] = ede_link("I", e)
        for c in (4, 5, 6):
            ws.cell(P + 8, c).alignment = Alignment(horizontal="center")
        merge(ws, f"D{P+9}:G{P+9}"); instr.append(f"D{P+9}"); phase.append(f"F{P+8}")
        a, b = outlet_rows(ws, rtu, 27, rtu, 29, P + 10, 6, "Register / Grille / Diffuser Airflow")
        total_row(ws, rtu, 43, P + 18, [(a, b)])
        remarks(ws, rtu, 52, P + 19, lines=1)
        ws[f"K{P+4}"] = f"=H{P+18}"; ws[f"L{P+4}"] = f"=L{P+18}"
        ws[f"M{P+4}"] = f'=IF(OR(L{P+4}="",K{P+4}="",K{P+4}=0),"",L{P+4}/K{P+4})'
        copy_style(fans["M9"], ws[f"M{P+4}"])
        for r in range(P, P + SF_BLOCK):
            if ws.row_dimensions[r].height is None:
                ws.row_dimensions[r].height = 13.35
        if u % 2 == 1:
            ws.row_breaks.append(Break(id=P + SF_BLOCK - 1))
    dv(ws, instr, "Airflow.Instrument")
    ws.print_area = f"A1:N{4 + 4 * SF_BLOCK - 1}"
    log("Small Fans: 40 blocks, two per page (NEBB 5.3.6 data set + 6 outlets each)")


# --------------------------------------------------------------------------- #
# 3. Traverses with duct point grids
# --------------------------------------------------------------------------- #
def traverses(wb):
    old = wb["Traverses"]
    ws = wb.create_sheet("Traverses (grid)")
    for col, dim in old.column_dimensions.items():
        if dim.width:
            ws.column_dimensions[col].width = dim.width
    for col in "BDEFGHIJKLM":
        ws.column_dimensions[col].width = 6.86
    ws.column_dimensions["C"].width = 7.5
    ws.sheet_format.defaultColWidth = 6.86
    copy_block(old, 1, 3, ws, 1)
    for m in old.merged_cells.ranges:
        if m.max_row <= 3:
            merge(ws, str(m))
    ws.print_title_rows = "1:3"; ws.page_margins = copy.copy(old.page_margins); ws.page_setup.scale = 97
    ws.sheet_view.showGridLines = old.sheet_view.showGridLines
    small = Font(name=old["B6"].font.name, size=8); smallb = Font(name=old["B6"].font.name, size=8, bold=True)
    grid_font = Font(name=old["B6"].font.name, size=8)
    shape_cells, instr_cells = [], []
    n = 0
    for pg in range(TRAV_PAGES):
        top = 4 + pg * 49
        for k in range(3):
            n += 1
            P = top + 1 + k * TRAV_BLOCK
            copy_block(old, 6, 9, ws, P)                       # title, labels, values, instrument rows
            ws[f"B{P+2}"] = f"T-{n}"
            r2, r3, r4, r5 = P + 2, P + 3, P + 4, P + 5
            # duct geometry row
            for c in range(2, 14):
                copy_style(old.cell(9, c), ws.cell(r4, c))
                ws.cell(r4, c).font = small
            for rng in (f"B{r4}:C{r4}", f"D{r4}:E{r4}"):
                merge(ws, rng)
            ws[f"B{r4}"] = "Duct shape:"; ws[f"B{r4}"].font = smallb
            ws[f"F{r4}"] = "W / Dia (in)"; ws[f"H{r4}"] = "H (in)"; ws[f"J{r4}"] = "Liner (in)"; ws[f"L{r4}"] = "Points"
            for c in ("F", "H", "J", "L"):
                ws[f"{c}{r4}"].font = smallb
            S, W, H, LN = f"$D${r4}", f"$G${r4}", f"$I${r4}", f"$K${r4}"
            nW = f'IF({S}="Round",IF({W}<=9,6,IF({W}<=12,8,10)),IF({W}<12,2,MIN(10,MAX(3,ROUNDUP({W}/6,0)))))'
            nH = f'IF({S}="Round",2,IF({H}<12,2,MIN(8,MAX(3,ROUNDUP({H}/6,0)))))'
            ws[f"M{r4}"] = f'=IF(OR({S}="",{W}=""),"",IF({S}="Round",{nW}&" x 2 axes",{nW}&" x "&{nH}))'
            ws[f"M{r4}"].font = small
            shape_cells.append(f"D{r4}")
            # size text and free area
            ws[f"F{r2}"] = f'=IF({W}="","",IF({S}="Round",{W}&""" dia",{W}&""" x "&{H}&""""))'
            ws[f"H{r2}"] = (f'=IF(OR({S}="",{W}=""),"",IF({S}="Round",ROUND(PI()*(({W}-2*N({LN}))/2)^2/144,3),'
                            f'IF({H}="","",ROUND(({W}-2*N({LN}))*({H}-2*N({LN}))/144,3))))')
            ws[f"H{r2}"].number_format = "0.000"
            # positions header row
            for c in range(2, 14):
                ws.cell(r5, c).font = smallb; ws.cell(r5, c).alignment = Alignment(horizontal="center")
                ws.cell(r5, c).border = Border(bottom=THIN); ws.cell(r5, c).fill = PatternFill("solid", fgColor="FFD9E2F3")
            ws[f"B{r5}"] = "Point"; ws[f"C{r5}"] = "Depth ↓ / Pos. → (in)"
            merge(ws, f"B{r5}:C{r5}")
            ws[f"B{r5}"] = "Depth↓ Pos→ (in)"
            for j in range(1, 11):
                col = get_column_letter(3 + j)
                pos_rect = f"ROUND(({j}-0.5)*{W}/{nW},1)"
                pos_round = (f"ROUND(IF({j}<={nW}/2,{W}/2*(1-SQRT(({nW}-2*{j}+1)/{nW})),"
                             f"{W}-{W}/2*(1-SQRT(({nW}-2*({nW}-{j}+1)+1)/{nW}))),1)")
                ws[f"{col}{r5}"] = f'=IF(OR({S}="",{W}="",{j}>{nW}),"",IF({S}="Round",{pos_round},{pos_rect}))'
            # grid rows
            g1, g2 = r5 + 1, r5 + 8
            for i in range(1, 9):
                r = r5 + i
                ws.row_dimensions[r].height = 13.35
                for c in range(2, 14):
                    cell = ws.cell(r, c); cell.font = grid_font; cell.alignment = Alignment(horizontal="center")
                    cell.border = Border(left=THIN, right=THIN, top=Side(style="hair"), bottom=Side(style="hair"))
                merge(ws, f"B{r}:C{r}")
                ws[f"B{r}"] = (f'=IF(OR({S}="",{W}=""),"",IF({S}="Round",IF({i}=1,"Axis 1 (0°)",IF({i}=2,"Axis 2 (90°)","")),'
                               f'IF({i}>{nH},"",ROUND(({i}-0.5)*{H}/{nH},1))))')
                ws[f"B{r}"].font = smallb; ws[f"B{r}"].fill = PatternFill("solid", fgColor="FFF2F2F2")
            # velocities: final = grid average; CFM = vel x free area
            ws[f"L{r2}"] = f'=IF(COUNT(D{g1}:M{g2})=0,"",ROUND(AVERAGE(D{g1}:M{g2}),0))'
            ws[f"K{r2}"] = f'=IF(N(J{r2})*N(H{r2})=0,"",ROUND(N(J{r2})*N(H{r2}),0))'
            ws[f"M{r2}"] = f'=IF(N(L{r2})*N(H{r2})=0,"",ROUND(N(L{r2})*N(H{r2}),0))'
            instr_cells.append(f"D{r3}")
            for r in range(P, P + TRAV_BLOCK):
                if ws.row_dimensions[r].height is None:
                    ws.row_dimensions[r].height = 13.35
        rr = top + 46
        remarks(ws, wb["RTUs"], 52, rr, lines=2)
        for r in range(rr, top + 49):
            ws.row_dimensions[r].height = 13.35
        ws.row_breaks.append(Break(id=top + 48))
    dv(ws, shape_cells, "Duct.Shape"); dv(ws, instr_cells, "Traverse.Instrument")
    ws["P5"] = ("Equal-area traverse per NEBB 6.3.3: rectangular <12\" per axis = 2 points, otherwise ≥3 points at ≤6\" spacing; "
                "round 6-9\" = 6 points, 10-12\" = 8, >12\" = 10 per axis on two axes at 90°; round ≤5\" may use 90% of the centerline "
                "velocity. Positions are measured from the duct wall (inside liner). Enter velocities (fpm) in the grid; Final VEL is the average.")
    ws["P5"].font = Font(size=8, italic=True)
    ws.print_area = f"A1:M{4 + 2 * 49 - 1}"
    ws.oddHeader.left.text = "&G"; ws.oddHeader.center.text = "Traverse Measurement Report"
    wb.remove(old); ws.title = "Traverses"
    dd = wb["{Dropdowns}"]
    dd["W1"] = "Duct Shape"; dd["W2"] = "Rectangular"; dd["W3"] = "Round"
    wb.defined_names["Duct.Shape"] = DefinedName("Duct.Shape", attr_text=f"{DD}!$W$2:$W$3")
    log(f"Traverses rebuilt: {TRAV_PAGES * 3} traverses, three per page, equal-area point grid with computed positions (NEBB 6.3.3)")


# --------------------------------------------------------------------------- #
# 4. Equipment Summary
# --------------------------------------------------------------------------- #
def equipment_summary(wb):
    bb = wb["Building Balance"]; rtu = wb["RTUs"]
    ws = wb.create_sheet("Equipment Summary")
    for col, w in zip("ABCDEFGHIJKLMN", (3.43, 9, 15, 6.6, 6.6, 5.2, 6.6, 6.6, 6.2, 6.2, 6.2, 6.2, 6.6, 3.14)):
        ws.column_dimensions[col].width = w
    copy_block(bb, 1, 3, ws, 1)
    for m in bb.merged_cells.ranges:
        if m.max_row <= 3:
            merge(ws, str(m))
    ws.print_title_rows = "1:3"; ws.page_margins = copy.copy(bb.page_margins); ws.page_setup.scale = 93
    ws.sheet_view.showGridLines = bb.sheet_view.showGridLines
    ws.oddHeader.left.text = "&G"; ws.oddHeader.center.text = "Equipment Summary"
    fname = bb["B6"].font.name
    hdr_fill = copy.copy(bb["B6"].fill); sec_fill = copy.copy(bb["B5"].fill)
    ws["B5"] = "NEBB / contract tolerance ±"; ws["B5"].font = Font(name=fname, size=9, bold=True); merge(ws, "B5:D5")
    ws["E5"] = 0.10; ws["E5"].number_format = "0%"; ws["E5"].font = Font(name=fname, size=9, bold=True, color="FF0000FF")
    ws["F5"] = "(Status = Check when |Actual/Design − 1| exceeds this)"; ws["F5"].font = Font(name=fname, size=8, italic=True)
    heads = ["Unit", "Area Served", "Design CFM", "Actual CFM", "%", "OA / Exh Des", "OA / Exh Act", "ESP Des", "ESP Act", "Fan RPM", "Amps", "Status"]
    r = 7
    for i, h in enumerate(heads):
        c = ws.cell(r, 2 + i); c.value = h; c.font = Font(name=fname, size=9, bold=True); c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.row_dimensions[r].height = 24
    row = 8
    TOL = "$E$5"

    def section(title):
        nonlocal row
        merge(ws, f"B{row}:M{row}"); ws[f"B{row}"] = title; ws[f"B{row}"].font = Font(name=fname, size=10, bold=True); ws[f"B{row}"].fill = sec_fill
        ws[f"B{row}"].alignment = Alignment(horizontal="left"); ws.row_dimensions[row].height = 13.35
        row += 1

    def line(unit, area, d, a, oa_d=None, oa_a=None, esp_d=None, esp_a=None, rpm=None, amps=None):
        nonlocal row
        vals = [unit, area, d, a, f'=IF(OR(E{row}="",D{row}="",D{row}=0),"",E{row}/D{row})', oa_d, oa_a, esp_d, esp_a, rpm, amps,
                f'=IF(OR(D{row}="",E{row}="",D{row}=0),"",IF(ABS(E{row}/D{row}-1)<={TOL},"OK","Check"))']
        for i, v in enumerate(vals):
            c = ws.cell(row, 2 + i); c.value = v; c.font = Font(name=fname, size=9)
            c.border = Border(bottom=Side(style="hair"), left=THIN if i == 0 else None, right=THIN if i == 11 else None)
            c.alignment = Alignment(horizontal="left" if i < 2 else "center")
        ws.cell(row, 6).number_format = "0%"; ws.cell(row, 13).font = Font(name=fname, size=9, bold=True)
        for col in (4, 5, 7, 8):
            ws.cell(row, col).number_format = "#,##0"
        ws.row_dimensions[row].height = 13.35
        row += 1

    def sref(sheet, cell):
        return f"=IF('{sheet}'!{cell}=\"\",\"\",'{sheet}'!{cell})"

    section("Rooftop Units")
    for u in range(40):
        P = 4 + u * 2 * PAGE
        line(sref("RTUs", f"D{P}"), sref("RTUs", f"I{P}"), sref("RTUs", f"K{P+5}"), sref("RTUs", f"L{P+5}"), sref("RTUs", f"K{P+6}"), sref("RTUs", f"L{P+6}"),
             sref("RTUs", f"K{P+8}"), sref("RTUs", f"L{P+8}"), sref("RTUs", f"L{P+18}"),
             f"=IF(COUNT('RTUs'!E{P+16}:G{P+16})=0,\"\",ROUND(AVERAGE('RTUs'!E{P+16}:G{P+16}),1))")
    section("Make-up Air Units")
    for u in range(10):
        P = 4 + u * 2 * PAGE
        line(sref("MAUs", f"D{P}"), sref("MAUs", f"I{P}"), sref("MAUs", f"K{P+5}"), sref("MAUs", f"L{P+5}"), None, None,
             sref("MAUs", f"K{P+6}"), sref("MAUs", f"L{P+6}"), sref("MAUs", f"L{P+18}"),
             f"=IF(COUNT('MAUs'!E{P+16}:G{P+16})=0,\"\",ROUND(AVERAGE('MAUs'!E{P+16}:G{P+16}),1))")
    section("Energy Recovery Units  (OA / Exh columns = exhaust airflow)")
    for u in range(10):
        P = 4 + u * 2 * PAGE
        line(sref("ERVs", f"D{P}"), sref("ERVs", f"I{P}"), sref("ERVs", f"K{P+5}"), sref("ERVs", f"L{P+5}"), sref("ERVs", f"K{P+6}"), sref("ERVs", f"L{P+6}"),
             None, None, sref("ERVs", f"L{P+18}"), f"=IF(COUNT('ERVs'!E{P+16}:G{P+16})=0,\"\",ROUND(AVERAGE('ERVs'!E{P+16}:G{P+16}),1))")
    section("Fans")
    for u in range(40):
        P = 4 + u * 2 * PAGE
        line(sref("Fans", f"D{P}"), sref("Fans", f"I{P}"), sref("Fans", f"K{P+5}"), sref("Fans", f"L{P+5}"), None, None,
             sref("Fans", f"K{P+6}"), sref("Fans", f"L{P+6}"), sref("Fans", f"L{P+18}"),
             f"=IF(COUNT('Fans'!E{P+16}:G{P+16})=0,\"\",ROUND(AVERAGE('Fans'!E{P+16}:G{P+16}),1))")
    section("Small Exhaust Fans")
    for u in range(N_SMALL):
        P = 4 + u * SF_BLOCK
        line(sref("Small Fans", f"D{P}"), sref("Small Fans", f"I{P}"), sref("Small Fans", f"K{P+4}"), sref("Small Fans", f"L{P+4}"), None, None,
             sref("Small Fans", f"K{P+5}"), sref("Small Fans", f"L{P+5}"), sref("Small Fans", f"L{P+6}"), sref("Small Fans", f"K{P+8}"))
    section("Kitchen Hoods  (Design / Actual = hood exhaust)")
    for a in anchors(wb["Hoods"]):
        line(sref("Hoods", f"D{a}"), sref("Hoods", f"I{a}"), sref("Hoods", f"B{a+18}"), sref("Hoods", f"F{a+18}"))
    ws.print_area = f"A1:N{row - 1}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    log(f"Equipment Summary: {row - 8} rows (RTU, MAU, ERV, fans, small fans, hoods) with tolerance flag")


# --------------------------------------------------------------------------- #
# 5. Building Balance: small fan rows on the exhaust side
# --------------------------------------------------------------------------- #
def building_balance(wb):
    ws = wb["Building Balance"]; NS = 20
    moved = [m for m in ws.merged_cells.ranges if m.min_row >= 67]
    for m in moved:
        ws.unmerge_cells(str(m))
    ws.move_range("B67:M84", rows=NS, cols=0, translate=False)
    for m in moved:
        merge(ws, f"{get_column_letter(m.min_col)}{m.min_row+NS}:{get_column_letter(m.max_col)}{m.max_row+NS}")
    for r in range(84, 66, -1):
        ws.row_dimensions[r + NS].height = ws.row_dimensions[r].height
    for i in range(NS):
        r = 67 + i; P = 4 + i * SF_BLOCK
        for c in range(2, 14):
            copy_style(ws.cell(66, c), ws.cell(r, c))
        for m in (f"C{r}:D{r}", f"E{r}:F{r}", f"I{r}:J{r}", f"K{r}:L{r}"):
            merge(ws, m)
        ws[f"H{r}"] = f"=IF('Small Fans'!D{P}=\"\",\"\",'Small Fans'!D{P})"
        ws[f"I{r}"] = f"=IF('Small Fans'!K{P+4}=\"\",\"\",'Small Fans'!K{P+4})"; ws[f"K{r}"] = f"=IF('Small Fans'!L{P+4}=\"\",\"\",'Small Fans'!L{P+4})"
        ws[f"M{r}"] = f'=IF(OR(K{r}="",I{r}="",I{r}=0),"",K{r}/I{r})'
        ws.row_dimensions[r].height = ws.row_dimensions[66].height
        ws.row_dimensions[r].hidden = i >= 2
    t = 67 + NS
    for col in "CEIK":
        ws[f"{col}{t}"] = f'=IF(SUM({col}7:{col}{t-1})=0,"",SUM({col}7:{col}{t-1}))'
    ws[f"G{t}"] = f'=IF(OR(E{t}="",C{t}=0),"",E{t}/C{t})'; ws[f"M{t}"] = f'=IF(OR(K{t}="",I{t}=0),"",K{t}/I{t})'
    ws[f"H{t+2}"] = f'=IF(AND(C{t}="",I{t}=""),"",N(C{t})-N(I{t}))'; ws[f"H{t+4}"] = f'=IF(AND(E{t}="",K{t}=""),"",N(E{t})-N(K{t}))'
    ws.print_area = f"A1:M{84 + NS}"
    log("Building Balance: 20 small-fan exhaust rows added; totals and balance re-based")


def toc_and_order(wb):
    toc = wb["ToC"]
    toc["C16"] = "Equipment Summary and Building Balance"
    toc["C28"] = "Fans, Small Exhaust Fans and VAV Terminals"
    order = ["Cover Page", "ToC", "Narrative", "Summary - New", "Summary - (E)", "{Project Information}", "{Equipment Data Entry}",
             "{Dropdowns}", "Equipment Summary", "Building Balance", "RTUs", "MAUs", "ERVs", "Fans", "Small Fans", "VAVs", "Hoods",
             "Traverses", "Photos", "Certification", "NEBB Cert ", "NEBB Frm Cert", "Abbreviations", "Calibration"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Cover Page"
    log("ToC updated; sheet order: " + ", ".join(order))


def build(out=OUT):
    wb = load_workbook(SRC, keep_vba=True)
    data_entry_small_fans(wb)
    small_fans(wb)
    traverses(wb)
    equipment_summary(wb)
    building_balance(wb)
    toc_and_order(wb)
    tmp = tempfile.mktemp(suffix=".xlsm")
    wb.save(tmp)
    xlsm_parts.restore(SRC, tmp, out, footer="&amp;C&amp;8Page &amp;P", skip_footer=("Cover Page",),
                       extra_sheet_sources={"Small Fans": "Fans", "Equipment Summary": "Building Balance"},
                       header_overrides={"Small Fans": "Small Exhaust Fan Report", "Equipment Summary": "Equipment Summary",
                                         "Traverses": "Traverse Measurement Report"})
    os.remove(tmp)
    with open(os.path.join(ROOT, "docs", "build-log-rev03.txt"), "w") as fh:
        fh.write("\n".join(LOG) + "\n")
    print("written", out)


if __name__ == "__main__":
    build()
