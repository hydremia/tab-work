"""Part 2 of the workbook build: new sheets, Evergreen hood method, NEBB fields."""
import copy
import re

from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

from build_workbook import EDE, DD, anchors, copy_style, log, setf

THIN = Side(style="thin")
EDE_RTU, EDE_MAU, EDE_ERV, EDE_FAN, EDE_HOOD, EDE_VAV = 7, 52, 65, 81, 126, 150
EDE_RE = re.compile(r"('\{Equipment Data Entry\}'!\$?[A-Z]+\$?)(\d+)")


def is_merged_slave(cell):
    return type(cell).__name__ == "MergedCell"


def merge(ws, rng):
    if rng not in [str(m) for m in ws.merged_cells.ranges]:
        ws.merge_cells(rng)


def unmerge_if(ws, pred):
    for m in list(ws.merged_cells.ranges):
        if pred(m):
            ws.unmerge_cells(str(m))


def ede_link(col, row, blank='""'):
    ref = f"{EDE}!{col}{row}"
    return f'=IF({ref}="",{blank},{ref})'


def label_font(ws, like):
    return copy.copy(ws[like].font)


# --------------------------------------------------------------------------- #
# generic: copy block-1 same-sheet formulas into blocks where they are missing
# --------------------------------------------------------------------------- #
def harmonize_formulas(wb):
    for name in ("RTU Data", "MAU Data", "Fan Data (EFs, TFs, etc.)"):
        ws = wb[name]; aa = anchors(ws); base = aa[0]; n = 0
        src = {}
        for row in ws.iter_rows(min_row=base, max_row=base + 22, max_col=13):
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and "!" not in c.value:
                    src[(c.row - base, c.column)] = c
        for anc in aa[1:]:
            for (off, col), s in src.items():
                dst = ws.cell(anc + off, col)
                if is_merged_slave(dst) or dst.value is not None:
                    continue
                dst.value = Translator(s.value, origin=s.coordinate).translate_formula(dst.coordinate)
                copy_style(s, dst); n += 1
        log(f"{name}: filled {n} missing block formulas from block 1 (corrected FLA etc.); block-1 same-sheet formulas={len(src)}")


# --------------------------------------------------------------------------- #
# Dropdown lists and Evergreen reference tables (hidden sheet)
# --------------------------------------------------------------------------- #
VELGRID_SIZES = [('10" x 16"', 0.78), ('10" x 20"', 0.99), ('12" x 12"', 0.69), ('12" x 16"', 0.97),
                 ('12" x 20"', 1.25), ('12" x 24"', 1.52), ('16" x 16"', 1.35), ('16" x 20"', 1.73),
                 ('16" x 25"', 2.22), ('20" x 20"', 2.23), ('20" x 25"', 2.85), ('24" x 24"', 3.36)]
SUPPLY_SIZES = [('12" x 12"', 0.69), ('12" x 16"', 0.97), ('12" x 20"', 1.25), ('12" x 24"', 1.52),
                ('16" x 16"', 1.36), ('16" x 20"', 1.75), ('16" x 25"', 2.24), ('20" x 20"', 2.25),
                ('20" x 25"', 2.88), ('24" x 24"', 3.36)]
CONDENSATE = [('10" x 16"', 0.131), ('10" x 20"', 0.196), ('12" x 16"', 0.131), ('12" x 20"', 0.196),
              ('16" x 16"', 0.131), ('16" x 20"', 0.196), ('20" x 16"', 0.262), ('20" x 20"', 0.349)]
HVC = [('16" Wide', 0.2402), ('20" Wide', 0.3027)]
FILTER_TYPES = ["Baffle (VelGrid)", "Captrate (VelGrid)", "Condensate Baffle (Airfoil)", "HVC / Slot (Airfoil)", "Supply Filter (VelGrid)"]
HOOD_INSTR = ["Evergreen VelGrid", "Evergreen Airfoil", "Other (see remarks)"]
AIR_INSTR = ["Flow Hood", "Velocity Grid", "Pitot Traverse", "Hot Wire Anemometer", "Rotating Vane Anemometer", "DDC / Controller Reading", "Other (see remarks)"]
DRIVE_TYPES = ["Belt", "Direct", "ECM"]
METHODS = ["Outlets", "PSP", "Filter Grid", "Profile Pressure", "Traverse"]
PSP = [(6, 0.88), (9, 0.88), (10, 0.88), (12, 0.88), (14, 0.95), (16, 0.95), (18, 0.95), (20, 0.95), (24, 0.95)]
PROFILE_P = [0.15, 0.20, 0.25, 0.30, 0.35, 0.40, 0.45, 0.50, 0.55, 0.60, 0.65]
PROFILE_CFM = {  # housing size -> CFM at each profile pressure (Evergreen "Direct Fired Profile Pres. CFM" rev 1)
    1: [697.15, 805.62, 1166.76, 1779.53, 2036.25, 2291.85, 2450.196, 2655.768, 2797.446, 2958.57, 3139.14],
    2: [2035.5, 2400, 2718.6, 3000, 3346.5, 4000, 5209.5, 5600, 5778.75, 6300, 6589.5],
    3: [1740.9, 1908.9, 2289, 4036.2, 4676.7, 4845, 4979.1, 5159.7, 6913.2, 7919.1, 8820],
    4: [3037.188, 3537.504, 3931.944, 4444.716, 4795.56, 6439.752, 6769.836, 6873.636, 7500, 13562.508, 16000],
    5: [4212.45, 9500, 12730.5, 15000, 17008.5, 19200, 20976, 21800, 22287, 23500, 24598.5],
}


def dropdown_tables(wb):
    ws = wb["{Dropdowns}"]
    hdr = Font(bold=True)
    ws["H1"], ws["I1"], ws["J1"], ws["K1"], ws["L1"] = "Key (Type|Size)", "Filter Size", "Free Area (sq ft)", "K-Factor", "Filter Type"
    r = 2
    rows = []
    for size, fa in VELGRID_SIZES:
        rows.append((FILTER_TYPES[0], size, fa, 1.28)); rows.append((FILTER_TYPES[1], size, fa, 1.34))
    rows += [(FILTER_TYPES[2], s, fa, 1.0) for s, fa in CONDENSATE]
    rows += [(FILTER_TYPES[3], s, fa, 1.0) for s, fa in HVC]
    rows += [(FILTER_TYPES[4], s, fa, 1.35) for s, fa in SUPPLY_SIZES]
    for t in FILTER_TYPES:
        rows.append((t, "No Filter", 0, 0))
    for t, s, fa, k in rows:
        ws[f"H{r}"] = f"{t}|{s}"; ws[f"I{r}"] = s; ws[f"J{r}"] = fa; ws[f"K{r}"] = k; ws[f"L{r}"] = t; r += 1
    ws["H1"].font = hdr
    ws["M1"] = "Filter Types"
    for i, t in enumerate(FILTER_TYPES):
        ws[f"M{i+2}"] = t
    sizes = ["No Filter"] + [s for s, _ in VELGRID_SIZES] + ['20" x 16"'] + [s for s, _ in HVC]
    ws["N1"] = "Filter Sizes"
    for i, s in enumerate(sizes):
        ws[f"N{i+2}"] = s
    ws["O1"] = "Hood Instruments"
    for i, s in enumerate(HOOD_INSTR):
        ws[f"O{i+2}"] = s
    ws["P1"] = "Airflow Instruments"
    for i, s in enumerate(AIR_INSTR):
        ws[f"P{i+2}"] = s
    ws["Q1"] = "Drive Types"
    for i, s in enumerate(DRIVE_TYPES):
        ws[f"Q{i+2}"] = s
    ws["R1"] = "PSP Width (in)"; ws["S1"] = "PSP K"
    for i, (w, k) in enumerate(PSP):
        ws[f"R{i+2}"] = w; ws[f"S{i+2}"] = k
    ws["T1"] = "Airflow Method"
    for i, s in enumerate(METHODS):
        ws[f"T{i+2}"] = s
    ws["U1"] = "Profile Pressure (in. w.g.)"
    for j in range(1, 6):
        ws.cell(1, 21 + j).value = f"Housing Size {j} CFM"
    for i, p in enumerate(PROFILE_P):
        ws[f"U{i+2}"] = p
        for j in range(1, 6):
            ws.cell(i + 2, 21 + j).value = PROFILE_CFM[j][i]
    ws["H45"] = ("Source: CaptiveAire / Evergreen Telemetry field worksheets (tb-worksheet-evergreen): Exhaust Baffle Filter rev 2.1, "
                 "Supply Fan Filter rev 1, Condensate Baffle rev 0, HVC/Slot rev 0, PSP Supply rev 2, Direct Fired Profile Pressure rev 1. "
                 "CFM = velocity x free area x K-factor. PSP CFM = avg velocity x (length - 2 - 2 x blanks) x width x K / 144.")
    for name, ref in (("Hood.FilterType", f"{DD}!$M$2:$M${1+len(FILTER_TYPES)}"),
                      ("Hood.FilterSize", f"{DD}!$N$2:$N${1+len(sizes)}"),
                      ("Hood.Instrument", f"{DD}!$O$2:$O${1+len(HOOD_INSTR)}"),
                      ("Airflow.Instrument", f"{DD}!$P$2:$P${1+len(AIR_INSTR)}"),
                      ("Drive.Type", f"{DD}!$Q$2:$Q${1+len(DRIVE_TYPES)}"),
                      ("PSP.Width", f"{DD}!$R$2:$R${1+len(PSP)}"),
                      ("Airflow.Method", f"{DD}!$T$2:$T${1+len(METHODS)}")):
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    log("{Dropdowns}: Evergreen filter/PSP/profile reference tables and new dropdown lists added")


def dv(ws, sqref, name):
    d = DataValidation(type="list", formula1=name, allow_blank=True)
    d.sqref = sqref
    ws.add_data_validation(d)


# --------------------------------------------------------------------------- #
# Equipment Data Entry: ERV columns, hood length, VAV section
# --------------------------------------------------------------------------- #
def data_entry_sections(wb):
    ws = wb["{Equipment Data Entry}"]
    setf(ws, "P64", "Design Supply CFM"); setf(ws, "Q64", "Design Exhaust CFM")
    ws["R64"] = "Design Supply ΔP"; ws["S64"] = "Design Exhaust ΔP"
    for c in ("R", "S"):
        copy_style(ws["Q64"], ws[f"{c}64"])
        for r in range(65, 79):
            copy_style(ws[f"Q{r}"], ws[f"{c}{r}"])
    ws["P63"] = "Linked to ERV Data"
    ws["B65"] = "ERV-1"
    ws["I125"] = "Hood Length (ft)"; copy_style(ws["H125"], ws["I125"])
    for r in range(126, 146):
        copy_style(ws[f"H{r}"], ws[f"I{r}"])
    # VAV section
    top = EDE_VAV - 2
    merge(ws, f"B{top}:O{top}")
    ws[f"B{top}"] = "Design Data VAV / Fan-Powered Terminals"
    copy_style(ws["B124"], ws[f"B{top}"])
    heads = ["Designation", "Area Served", "Location", "Manufacturer", "Model Number", "Inlet Size", "Terminal Type",
             "Design Max CFM", "Design Min CFM", "Heating CFM", "Fan CFM", "DDC Address"]
    for i, h in enumerate(heads):
        c = ws.cell(top + 1, 2 + i); c.value = h; copy_style(ws.cell(125, 2 + min(i, 6)), c)
    for r in range(EDE_VAV, EDE_VAV + 80):
        for i in range(len(heads)):
            copy_style(ws.cell(126, 2 + min(i, 6)), ws.cell(r, 2 + i))
        ws.row_dimensions[r].height = ws.row_dimensions[126].height
    ws[f"B{EDE_VAV}"] = "VAV-1"
    ws.row_dimensions[top + 1].height = ws.row_dimensions[125].height
    ws.print_area = f"A1:S{EDE_VAV + 80}"
    log("{Equipment Data Entry}: ERV ΔP columns, hood length column and an 80-row VAV terminal section added")


# --------------------------------------------------------------------------- #
# ERV sheets (cloned from MAU sheets)
# --------------------------------------------------------------------------- #
def clone_sheet(wb, src_name, new_name, header):
    src = wb[src_name]
    ws = wb.copy_worksheet(src)
    ws.title = new_name
    ws.sheet_properties.codeName = None  # never duplicate a VBA code name
    ws.print_area = src.print_area.split("!")[-1] if src.print_area else None
    ws.print_title_rows = src.print_title_rows
    for b in src.row_breaks.brk:
        ws.row_breaks.append(Break(id=b.id))
    for rng, rules in src.conditional_formatting._cf_rules.items():
        for rule in rules:
            ws.conditional_formatting.add(str(rng.sqref), copy.copy(rule))
    for d in src.data_validations.dataValidation:
        nd = DataValidation(type=d.type, formula1=d.formula1, allow_blank=d.allow_blank); nd.sqref = copy.copy(d.sqref)
        ws.add_data_validation(nd)
    ws.oddHeader.left.text = "&G"
    ws.oddHeader.center.text = header; ws.oddHeader.center.font = "+,Bold"; ws.oddHeader.center.size = 16
    ws.sheet_view.zoomScale = src.sheet_view.zoomScale
    return ws


def shift_ede_rows(ws, delta, only_rows=None):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "{Equipment Data Entry}" in c.value:
                def rep(m):
                    r = int(m.group(2))
                    return f"{m.group(1)}{r + delta}" if (only_rows is None or r in only_rows) else m.group(0)
                c.value = EDE_RE.sub(rep, c.value)


def erv_sheets(wb):
    data = clone_sheet(wb, "MAU Data", "ERV Data", "Energy Recovery Unit Data Report")
    air = clone_sheet(wb, "MAU Airflow", "ERV Airflow", "Energy Recovery Unit Airflow Measurement Report")
    shift_ede_rows(data, EDE_ERV - EDE_MAU)
    for row in data.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "'MAU Airflow'!" in c.value:
                c.value = c.value.replace("'MAU Airflow'!", "'ERV Airflow'!")
    aa = anchors(air); da = anchors(data)
    # ERV Airflow: block pairs per unit (supply outlets, exhaust inlets)
    for b, anc in enumerate(aa):
        u = b // 2
        air[f"D{anc}"] = ede_link("B", EDE_ERV + u); air[f"I{anc}"] = ede_link("C", EDE_ERV + u)
        air[f"B{anc+3}"] = "Supply Air Outlet Airflow" if b % 2 == 0 else "Exhaust Air Inlet Airflow"
    # ERV Data block performance rows
    for i, anc in enumerate(da):
        sup, exh = aa[2 * i] + 19, aa[2 * i + 1] + 19
        r9, r10, r11, r12, r13 = anc + 5, anc + 6, anc + 7, anc + 8, anc + 9
        for r in (r12, r13):
            for c in range(9, 14):
                copy_style(data.cell(r11, c), data.cell(r, c))
            merge(data, f"I{r}:J{r}")
        data[f"I{r9}"] = "Supply Airflow"; data[f"K{r9}"] = f"='ERV Airflow'!H{sup}"; data[f"L{r9}"] = f"='ERV Airflow'!L{sup}"
        data[f"M{r9}"] = f'=IF(OR(L{r9}="",K{r9}="",K{r9}=0),"",L{r9}/K{r9})'
        data[f"I{r10}"] = "Exhaust Airflow"; data[f"K{r10}"] = f"='ERV Airflow'!H{exh}"; data[f"L{r10}"] = f"='ERV Airflow'!L{exh}"
        data[f"M{r10}"] = f'=IF(OR(L{r10}="",K{r10}="",K{r10}=0),"",L{r10}/K{r10})'
        copy_style(data[f"M{r9}"], data[f"M{r10}"])
        data[f"I{r11}"] = "Supply ΔP (core)"; data[f"K{r11}"] = ede_link("R", EDE_ERV + i); data[f"L{r11}"] = None
        data[f"I{r12}"] = "Exhaust ΔP (core)"; data[f"K{r12}"] = ede_link("S", EDE_ERV + i); data[f"L{r12}"] = None
        data[f"I{r13}"] = "Fan RPMs"; data[f"K{r13}"] = ede_link("I", EDE_ERV + i); data[f"L{r13}"] = f'=IF(L{anc+18}="","",L{anc+18})'
        data[f"M{r11}"] = None; data[f"M{r12}"] = None
    log("ERV Data / ERV Airflow sheets created (NEBB 5.3.25: supply and exhaust airflow and ΔP), linked to Data Entry rows 65-74")
    return data, air


# --------------------------------------------------------------------------- #
# VAV sheets linked to Data Entry
# --------------------------------------------------------------------------- #
def vav_links(wb):
    ws = wb["VAV Data"]; aa = anchors(ws)
    for i, a in enumerate(aa):
        r = EDE_VAV + i
        ws[f"D{a}"] = ede_link("B", r); ws[f"I{a}"] = ede_link("C", r)
        ws[f"D{a+4}"] = ede_link("E", r); ws[f"D{a+5}"] = ede_link("F", r)
        ws[f"D{a+7}"] = ede_link("C", r); ws[f"D{a+8}"] = ede_link("D", r); ws[f"D{a+9}"] = ede_link("G", r)
        ws[f"B{a+9}"] = "Inlet Size"
        ws[f"L{a+5}"] = ede_link("I", r); ws[f"L{a+6}"] = ede_link("J", r); ws[f"L{a+7}"] = ede_link("L", r)
        ws[f"I{a+8}"] = "DDC Address"; ws[f"L{a+8}"] = ede_link("M", r)
        # row a+2: terminal type / instrument
        for c in range(2, 14):
            copy_style(ws.cell(a + 4, c), ws.cell(a + 2, c))
        merge(ws, f"B{a+2}:C{a+2}"); merge(ws, f"D{a+2}:G{a+2}"); merge(ws, f"I{a+2}:K{a+2}"); merge(ws, f"L{a+2}:M{a+2}")
        ws[f"B{a+2}"] = "Terminal Type"; ws[f"D{a+2}"] = ede_link("H", r); ws[f"I{a+2}"] = "Instrument"
        ws.row_dimensions[a + 2].height = ws.row_dimensions[a + 4].height
        # row a+10: DDC max/min, heating CFM
        for c in range(2, 14):
            copy_style(ws.cell(a + 9, c), ws.cell(a + 10, c))
        merge(ws, f"B{a+10}:C{a+10}"); merge(ws, f"D{a+10}:G{a+10}"); merge(ws, f"I{a+10}:K{a+10}")
        ws[f"B{a+10}"] = "DDC Max / Min"; ws[f"I{a+10}"] = "Heating CFM"; ws[f"L{a+10}"] = ede_link("K", r)
        ws.row_dimensions[a + 10].height = ws.row_dimensions[a + 9].height
    dv(ws, " ".join(f"L{a+2}" for a in aa), "Airflow.Instrument")
    air = wb["VAV 1-20 Airflow"]
    for i, a in enumerate(anchors(air)):
        air[f"D{a}"] = ede_link("B", EDE_VAV + i); air[f"I{a}"] = ede_link("C", EDE_VAV + i)
    log("VAV Data / VAV 1-20 Airflow: linked to Data Entry VAV section; terminal type, instrument, DDC max/min and heating CFM added (NEBB 5.3.8-5.3.11)")


# --------------------------------------------------------------------------- #
# Hoods: Evergreen method
# --------------------------------------------------------------------------- #
def hoods_rebuild(wb):
    ws = wb["Hoods"]; aa = anchors(ws)
    ws.data_validations.dataValidation = [d for d in ws.data_validations.dataValidation if "Filter_Size" not in str(d.formula1)]
    unmerge_if(ws, lambda m: m.min_col == 16 and m.max_col == 26)  # P:Z technician-note merges
    type_cells, instr_cells, size_cells = [], [], []
    for i, a in enumerate(aa):
        ft, ins = a + 14, a + 15
        # left column
        for c in range(2, 8):
            copy_style(ws.cell(a + 6, c), ws.cell(a + 7, c)); copy_style(ws.cell(a + 13, c), ws.cell(ins, c))
        merge(ws, f"B{a+7}:D{a+7}"); merge(ws, f"E{a+7}:G{a+7}")
        merge(ws, f"B{ins}:D{ins}"); merge(ws, f"E{ins}:G{ins}")
        ws[f"B{a+7}"] = "Hood Length (ft)"; ws[f"E{a+7}"] = ede_link("I", EDE_HOOD + i)
        ws[f"B{ft}"] = "Filter Type"; ws[f"E{ft}"] = None
        ws[f"B{ins}"] = "Instrument"; ws[f"E{ins}"] = None
        type_cells.append(f"E{ft}"); instr_cells.append(f"E{ins}")
        # right table
        ws[f"I{a+4}"] = "Filter Size"
        ws[f"J{a+4}"] = "Initial VEL"; ws[f"L{a+4}"] = "Final VEL"
        for r in range(a + 6, a + 20):
            key = f'$E${ft}&"|"&$I{r}'
            fa = f"IFERROR(INDEX({DD}!$J:$J,MATCH({key},{DD}!$H:$H,0)),0)"
            kf = f"IFERROR(INDEX({DD}!$K:$K,MATCH({key},{DD}!$H:$H,0)),0)"
            ws[f"K{r}"] = f'=IF(OR($I{r}="",J{r}=""),"",J{r}*{fa}*{kf})'
            ws[f"M{r}"] = f'=IF(OR($I{r}="",L{r}=""),"",L{r}*{fa}*{kf})'
            ws[f"J{r}"] = f'=IF(COUNT(P{r}:R{r})=0,"",AVERAGE(P{r}:R{r}))'
            ws[f"L{r}"] = f'=IF(COUNT(S{r}:U{r})=0,"",AVERAGE(S{r}:U{r}))'
            size_cells.append(f"I{r}")
        hdr = a + 5
        for col, txt in zip("PQRSTU", ("Init 1", "Init 2", "Init 3", "Final 1", "Final 2", "Final 3")):
            ws[f"{col}{hdr}"] = txt; ws[f"{col}{hdr}"].font = Font(size=8, bold=True)
        ws[f"P{a+4}"] = "Velocity readings (fpm) – 3 per filter for Airfoil; J/L average them (or type a single value over the formula)"
        ws[f"P{a+4}"].font = Font(size=8, italic=True)
        ws[f"P{a}"] = "Technician Notes"
        # CFM per foot row
        cf = a + 19
        for c in range(2, 8):
            copy_style(ws.cell(a + 18, c), ws.cell(cf, c))
        merge(ws, f"B{cf}:C{cf}"); merge(ws, f"D{cf}:E{cf}"); merge(ws, f"F{cf}:G{cf}")
        ws[f"B{cf}"] = "CFM / ft"
        ws[f"D{cf}"] = f'=IF(OR(D{a+18}="",E{a+7}="",E{a+7}=0),"",D{a+18}/E{a+7})'
        ws[f"F{cf}"] = f'=IF(OR(F{a+18}="",E{a+7}="",E{a+7}=0),"",F{a+18}/E{a+7})'
        ws[f"D{cf}"].number_format = "0"; ws[f"F{cf}"].number_format = "0"
    dv(ws, " ".join(type_cells), "Hood.FilterType")
    dv(ws, " ".join(instr_cells), "Hood.Instrument")
    dv(ws, " ".join(size_cells), "Hood.FilterSize")
    log("Hoods: CFM now = velocity x Evergreen free area x K-factor by filter type/size; filter type, instrument, hood length, CFM/ft and 3-reading averaging added")


# --------------------------------------------------------------------------- #
# MAU Supply Methods sheet (PSP, filter grid, direct-fired profile pressure)
# --------------------------------------------------------------------------- #
def mau_methods(wb):
    src = wb["MAU Airflow"]
    ws = wb.create_sheet("MAU Supply Methods")
    for col, dim in src.column_dimensions.items():
        ws.column_dimensions[col].width = dim.width
    for r in range(1, 4):
        ws.row_dimensions[r].height = src.row_dimensions[r].height
        for c in range(1, 15):
            copy_style(src.cell(r, c), ws.cell(r, c))
            if src.cell(r, c).value is not None:
                ws.cell(r, c).value = src.cell(r, c).value
    for m in src.merged_cells.ranges:
        if m.max_row <= 3:
            merge(ws, str(m))
    ws.oddHeader.left.text = "&G"; ws.oddHeader.center.text = "Make-up Air Supply Measurement Methods"
    ws.oddHeader.center.font = "+,Bold"; ws.oddHeader.center.size = 16
    ws.print_title_rows = "1:3"
    ws.page_margins = copy.copy(src.page_margins); ws.page_setup.scale = src.page_setup.scale
    ws.page_setup.orientation = src.page_setup.orientation
    lab = src["B8"]; hdr = src["B7"]; inp = src["I9"]; sys_lab = src["B4"]
    H = 24  # block height
    a0 = 4
    ws.sheet_properties.pageSetUpPr = copy.copy(src.sheet_properties.pageSetUpPr)
    method_cells = []
    for u in range(10):
        a = a0 + u * H
        # system row
        for c in range(2, 14):
            copy_style(src.cell(4, c), ws.cell(a, c)); copy_style(src.cell(5, c), ws.cell(a + 1, c))
        merge(ws, f"B{a}:C{a+1}"); merge(ws, f"D{a}:F{a+1}"); merge(ws, f"G{a}:H{a+1}"); merge(ws, f"I{a}:M{a+1}")
        ws[f"B{a}"] = "System"; ws[f"D{a}"] = ede_link("B", EDE_MAU + u); ws[f"G{a}"] = "Service"; ws[f"I{a}"] = ede_link("C", EDE_MAU + u)

        def section(r, title):
            for c in range(2, 14):
                copy_style(hdr, ws.cell(r, c))
            merge(ws, f"B{r}:M{r}"); ws[f"B{r}"] = title

        def field(r, c1, c2, text, v1, v2=None, formula=None, fmt=None):
            for c in range(c1, c2 + 1):
                copy_style(lab if c < v1 else inp, ws.cell(r, c))
            if c1 < v1 - 1:
                merge(ws, f"{get_column_letter(c1)}{r}:{get_column_letter(v1-1)}{r}")
            if v2 and v2 > v1:
                merge(ws, f"{get_column_letter(v1)}{r}:{get_column_letter(v2)}{r}")
            ws.cell(r, c1).value = text
            if formula:
                ws.cell(r, v1).value = formula
            if fmt:
                ws.cell(r, v1).number_format = fmt
        # --- PSP
        section(a + 3, "Perforated Supply Plenum (Evergreen VelGrid) – CFM = avg of readings × (L − 2 − 2×blanks) × W × K / 144")
        field(a + 4, 2, 4, "Length (in)", 4); field(a + 4, 5, 7, "Width (in)", 7); field(a + 4, 8, 10, "Blanks", 10)
        field(a + 4, 11, 13, "K", 13, formula=f'=IF(G{a+4}="","",IFERROR(INDEX({DD}!$S:$S,MATCH(G{a+4},{DD}!$R:$R,0)),""))', fmt="0.00")
        for k, r in enumerate((a + 5, a + 6)):
            field(r, 2, 13, "Readings 1-10 (fpm)" if k == 0 else "Readings 11-20 (fpm)", 4, 13)
            ws.unmerge_cells(f"D{r}:M{r}")
            for c in range(4, 14):
                copy_style(inp, ws.cell(r, c))
            merge(ws, f"B{r}:C{r}")
        rd = f"D{a+5}:M{a+6}"
        field(a + 7, 2, 7, "PSP CFM", 5, 7, formula=f'=IF(OR(D{a+4}="",G{a+4}="",COUNT({rd})=0),"",AVERAGE({rd})*(D{a+4}-2-2*N(J{a+4}))*G{a+4}*N(M{a+4})/144)', fmt="0")
        field(a + 7, 8, 13, "CFM / ft", 11, 13, formula=f'=IF(OR(E{a+7}="",D{a+4}=""),"",E{a+7}/(D{a+4}/12))', fmt="0")
        # --- filter grid
        section(a + 9, "Supply Filter Grid (Evergreen VelGrid, K = 1.35) – CFM = velocity × free area × K")
        field(a + 10, 2, 13, "Filter Size", 3, 13); ws.unmerge_cells(f"C{a+10}:M{a+10}")
        field(a + 11, 2, 13, "Velocity (fpm)", 3, 13); ws.unmerge_cells(f"C{a+11}:M{a+11}")
        field(a + 12, 2, 13, "CFM", 3, 13); ws.unmerge_cells(f"C{a+12}:M{a+12}")
        for c in range(3, 14):
            L = get_column_letter(c)
            for r in (a + 10, a + 11, a + 12):
                copy_style(inp, ws.cell(r, c))
            key = f'"Supply Filter (VelGrid)|"&{L}{a+10}'
            ws[f"{L}{a+12}"] = (f'=IF(OR({L}{a+10}="",{L}{a+11}=""),"",{L}{a+11}*IFERROR(INDEX({DD}!$J:$J,MATCH({key},{DD}!$H:$H,0)),0)'
                                f'*IFERROR(INDEX({DD}!$K:$K,MATCH({key},{DD}!$H:$H,0)),0))')
            ws[f"{L}{a+12}"].number_format = "0"
        field(a + 13, 2, 7, "Filter Grid Total CFM", 5, 7, formula=f'=IF(SUM(C{a+12}:M{a+12})=0,"",SUM(C{a+12}:M{a+12}))', fmt="0")
        # --- profile pressure
        section(a + 15, "Direct-Fired Heater – Supply CFM from Burner Profile Pressure (Evergreen table, linear interpolation)")
        field(a + 16, 2, 4, "Housing Size (1-5)", 4); field(a + 16, 5, 8, "Profile Pressure (in. w.g.)", 8, fmt="0.00")
        P, S = f"H{a+16}", f"D{a+16}"
        prs = f"{DD}!$U$2:$U$12"; tbl = f"{DD}!$V$2:$Z$12"
        idx = f"MATCH({P},{prs},1)"
        interp = (f"INDEX({tbl},{idx},{S})+({P}-INDEX({prs},{idx}))*"
                  f"(INDEX({tbl},MIN({idx}+1,11),{S})-INDEX({tbl},{idx},{S}))/0.05")
        field(a + 16, 9, 13, "Derived CFM", 11, 13,
              formula=f'=IF(OR({S}="",{P}=""),"",IF({P}<0.15,"Pressure too low",IF({P}>0.65,"Pressure too high",{interp})))', fmt="0")
        # --- method selection
        section(a + 18, "Airflow Basis for MAU Data Page")
        field(a + 19, 2, 7, "Method used", 5, 7)
        field(a + 19, 8, 13, "Design CFM (if no outlets)", 11, 13)
        field(a + 20, 2, 7, "Method Total CFM", 5, 7,
              formula=(f'=IF(E{a+19}="","",IF(E{a+19}="PSP",E{a+7},IF(E{a+19}="Filter Grid",E{a+13},'
                       f'IF(E{a+19}="Profile Pressure",K{a+16},""))))'), fmt="0")
        field(a + 20, 8, 13, "Remarks", 11, 13)
        method_cells.append((f"E{a+19}", f"K{a+19}", f"E{a+20}"))
        for r in range(a, a + H):
            ws.row_dimensions[r].height = 13.35
    ws.print_area = f"A1:N{a0 + 2 * H - 1}"
    ws.row_breaks.append(Break(id=a0 + 2 * H - 1))
    for k in range(2, 5):
        ws.row_breaks.append(Break(id=a0 + 2 * k * H - 1))
    dv(ws, " ".join(f"G{a0+u*H+4}" for u in range(10)), "PSP.Width")
    dv(ws, " ".join(f"{get_column_letter(c)}{a0+u*H+10}" for u in range(10) for c in range(3, 14)), "Hood.FilterSize")
    dv(ws, " ".join(f"E{a0+u*H+19}" for u in range(10)), "Airflow.Method")
    d = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    d.sqref = " ".join(f"D{a0+u*H+16}" for u in range(10)); ws.add_data_validation(d)
    # MAU Data: total airflow uses the method total when a method is selected
    md = wb["MAU Data"]
    for u, anc in enumerate(anchors(md)):
        m_sel, m_des, m_tot = method_cells[u]
        md[f"K{anc+5}"] = f"=IF('MAU Supply Methods'!{m_des}=\"\",{md[f'K{anc+5}'].value[1:]},'MAU Supply Methods'!{m_des})"
        md[f"L{anc+5}"] = f"=IF(OR('MAU Supply Methods'!{m_sel}=\"\",'MAU Supply Methods'!{m_sel}=\"Outlets\"),{md[f'L{anc+5}'].value[1:]},'MAU Supply Methods'!{m_tot})"
    log("MAU Supply Methods sheet added (PSP, supply filter grid, direct-fired profile pressure, Evergreen constants); MAU Data total airflow follows the selected method")
    return ws


# --------------------------------------------------------------------------- #
# NEBB fields on Data blocks, Airflow sheets, Traverses
# --------------------------------------------------------------------------- #
def data_block_fields(wb):
    for name in ("RTU Data", "MAU Data", "Fan Data (EFs, TFs, etc.)", "ERV Data"):
        ws = wb[name]; aa = anchors(ws)
        drive_cells = []
        for a in aa:
            r6, r14, r19 = a + 2, a + 10, a + 15
            lf = Font(name=ws[f"B{a+4}"].font.name, size=8, bold=True)
            vf = Font(name=ws[f"B{a+4}"].font.name, size=8)
            ws[f"B{r6}"] = "Drive Type:"; ws[f"B{r6}"].font = lf; ws[f"D{r6}"].font = vf; ws[f"D{r6}"].border = Border(bottom=THIN)
            ws[f"E{r6}"] = "Fan Rotation – Design:"; ws[f"E{r6}"].font = lf; ws[f"G{r6}"].font = vf; ws[f"G{r6}"].border = Border(bottom=THIN)
            ws[f"I{r6}"] = "Actual:"; ws[f"I{r6}"].font = lf; ws[f"J{r6}"].font = vf; ws[f"J{r6}"].border = Border(bottom=THIN)
            ws[f"K{r6}"] = "Sheave Bore (M/F):"; ws[f"K{r6}"].font = lf; ws[f"M{r6}"].font = vf; ws[f"M{r6}"].border = Border(bottom=THIN)
            ws[f"D{r6}"].alignment = Alignment(horizontal="center"); ws[f"G{r6}"].alignment = Alignment(horizontal="center")
            ws[f"J{r6}"].alignment = Alignment(horizontal="center"); ws[f"M{r6}"].alignment = Alignment(horizontal="center")
            drive_cells.append(f"D{r6}")
            for r, text in ((r14, "Filters: Type / Size / Qty"), (r19, "Final Settings / Setpoints")):
                for c in range(9, 14):
                    copy_style(ws.cell(a + 9, c), ws.cell(r, c))
                merge(ws, f"I{r}:J{r}"); merge(ws, f"K{r}:M{r}")
                ws[f"I{r}"] = text; ws[f"K{r}"] = None
            setf(ws, f"I{a+12}", "Motor Shv OD/Bore/PD"); setf(ws, f"I{a+13}", "Fan Shv OD/Bore")
        dv(ws, " ".join(drive_cells), "Drive.Type")
        log(f"{name}: drive type, fan rotation, sheave bore, filter data and final-settings fields added inside the printed block (NEBB 5.3.1-5.3.5)")


def airflow_sheet_fields(wb):
    for name in ("RTU Airflow", "MAU Airflow", "Fan Airflow", "VAV 1-20 Airflow", "ERV Airflow"):
        ws = wb[name]; aa = anchors(ws)
        # Type column: split the merged Area Served (C:D)
        unmerge_if(ws, lambda m: m.min_col == 3 and m.max_col == 4 and m.min_row == m.max_row and ws.cell(m.min_row, 3).value != "Total")
        cw = ws.column_dimensions["C"].width or 8.43; dw = ws.column_dimensions["D"].width or 8.43
        ws.column_dimensions["C"].width = round(cw + dw - 5.4, 2); ws.column_dimensions["D"].width = 5.4
        n = 0
        for row in ws.iter_rows(min_col=3, max_col=3):
            c = row[0]
            if c.value == "Area Served":
                d = ws.cell(c.row, 4); d.value = "Type"; copy_style(c, d); n += 1
        instr = []
        for a in aa:
            r = a + 2
            ws[f"B{r}"] = "Instrument:"; ws[f"B{r}"].font = Font(name=ws[f"B{a+4}"].font.name, size=8, bold=True)
            merge(ws, f"D{r}:F{r}"); ws[f"D{r}"].border = Border(bottom=THIN); ws[f"D{r}"].font = Font(name=ws[f"B{a+4}"].font.name, size=8)
            ws[f"D{r}"].alignment = Alignment(horizontal="center")
            ws[f"G{r}"] = "Ak basis / notes:"; ws[f"G{r}"].font = Font(name=ws[f"B{a+4}"].font.name, size=8, bold=True)
            merge(ws, f"I{r}:M{r}"); ws[f"I{r}"].border = Border(bottom=THIN); ws[f"I{r}"].font = Font(name=ws[f"B{a+4}"].font.name, size=8)
            instr.append(f"D{r}")
        dv(ws, " ".join(instr), "Airflow.Instrument")
        log(f"{name}: air-device Type column ({n} tables) and per-system Instrument / Ak-basis line added (NEBB 5.3.7)")


def traverse_profiles(wb):
    ws = wb["Traverses"]
    tops = anchors(ws, label="Airflow Traverse Measurement")
    for t in tops:
        d, p1, p2 = t + 2, t + 4, t + 5
        for r in (p1, p2):
            ws.row_dimensions[r].height = 13.35
            for c in range(3, 14):
                cell = ws.cell(r, c); cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
                cell.font = Font(name=ws[f"B{t+1}"].font.name, size=8); cell.alignment = Alignment(horizontal="center")
        merge(ws, f"B{p1}:B{p2}")
        ws[f"B{p1}"] = f'="Profile ("&COUNT(C{p1}:M{p2})&" pts)"'
        ws[f"B{p1}"].font = Font(name=ws[f"B{t+1}"].font.name, size=7, bold=True)
        ws[f"B{p1}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws[f"B{p1}"].border = Border(left=THIN, top=THIN, bottom=THIN); ws[f"B{p2}"].border = Border(left=THIN, bottom=THIN)
        ws[f"L{d}"] = f'=IF(COUNT(C{p1}:M{p2})=0,"",ROUND(AVERAGE(C{p1}:M{p2}),0))'
    for r in range(3, 7):
        ws[f"P{r}"] = None
    ws["P6"] = "Velocity profile: enter each traverse reading (fpm); Final VEL averages them. Overflow readings may be placed in P:AA and included by editing the formula."
    ws["P6"].font = Font(size=8, italic=True)
    log(f"Traverses: 22-point velocity profile grid and reading count added to all {len(tops)} traverse blocks; Final VEL averages the profile (NEBB 5.3.12)")


# --------------------------------------------------------------------------- #
# Building Balance: ERV rows
# --------------------------------------------------------------------------- #
def building_balance_erv(wb):
    ws = wb["Building Balance"]; erv = wb["ERV Data"]; ea = anchors(erv)
    moved = [m for m in ws.merged_cells.ranges if m.min_row >= 57]
    for m in moved:
        ws.unmerge_cells(str(m))
    ws.move_range("B57:M74", rows=10, cols=0, translate=True)
    for m in moved:
        merge(ws, f"{get_column_letter(m.min_col)}{m.min_row+10}:{get_column_letter(m.max_col)}{m.max_row+10}")
    for r in range(74, 56, -1):
        ws.row_dimensions[r + 10].height = ws.row_dimensions[r].height
    for i in range(10):
        r = 57 + i; anc = ea[i]
        for c in range(2, 8):
            copy_style(ws.cell(47, c), ws.cell(r, c))
        for c in range(8, 14):
            copy_style(ws.cell(7, c), ws.cell(r, c))
        for m in (f"C{r}:D{r}", f"E{r}:F{r}", f"I{r}:J{r}", f"K{r}:L{r}"):
            merge(ws, m)
        ws[f"B{r}"] = ede_link("B", EDE_ERV + i)
        ws[f"C{r}"] = f"=IF('ERV Data'!K{anc+5}=\"\",\"\",'ERV Data'!K{anc+5})"; ws[f"E{r}"] = f"=IF('ERV Data'!L{anc+5}=\"\",\"\",'ERV Data'!L{anc+5})"
        ws[f"G{r}"] = f'=IF(OR(E{r}="",C{r}="",C{r}=0),"",E{r}/C{r})'
        ws[f"H{r}"] = ede_link("B", EDE_ERV + i)
        ws[f"I{r}"] = f"=IF('ERV Data'!K{anc+6}=\"\",\"\",'ERV Data'!K{anc+6})"; ws[f"K{r}"] = f"=IF('ERV Data'!L{anc+6}=\"\",\"\",'ERV Data'!L{anc+6})"
        ws[f"M{r}"] = f'=IF(OR(K{r}="",I{r}="",I{r}=0),"",K{r}/I{r})'
        ws.row_dimensions[r].height = ws.row_dimensions[47].height
        ws.row_dimensions[r].hidden = i >= 2
    for col in "CEIK":
        ws[f"{col}67"] = f'=IF(SUM({col}7:{col}66)=0,"",SUM({col}7:{col}66))'
    ws["G67"] = '=IF(OR(E67="",C67=0),"",E67/C67)'; ws["M67"] = '=IF(OR(K67="",I67=0),"",K67/I67)'
    ws["H69"] = '=IF(AND(C67="",I67=""),"",N(C67)-N(I67))'; ws["H71"] = '=IF(AND(E67="",K67=""),"",N(E67)-N(K67))'
    # pressure table replaces the single measured-pressure cell and the old Notes block
    for m in [m for m in ws.merged_cells.ranges if m.min_row >= 73]:
        ws.unmerge_cells(str(m))
    for row in ws.iter_rows(min_row=73, max_row=90):
        for c in row:
            c.value = None
    merge(ws, "E73:G74"); ws["E73"] = "Measured Building Pressures (in. w.g.)"
    rows = [("Building", "Outdoors"), ("Kitchen", "Dining"), ("", "")]
    hdr_font = Font(name=ws["E73"].font.name, size=9, bold=True)
    ws["B76"], ws["E76"], ws["H76"], ws["K76"] = "Test Space", "Reference Space", "ΔP (in. w.g.)", "Remarks"
    for c in ("B76", "E76", "H76", "K76"):
        ws[c].font = hdr_font; ws[c].border = Border(bottom=THIN)
    for i, (ts, rs) in enumerate(rows):
        r = 77 + i
        ws[f"B{r}"] = ts; ws[f"E{r}"] = rs
        for c in (2, 5, 8, 11):
            ws.cell(r, c).border = Border(bottom=Side(style="hair")); ws.cell(r, c).font = Font(name=ws["E73"].font.name, size=9)
        merge(ws, f"B{r}:D{r}"); merge(ws, f"E{r}:G{r}"); merge(ws, f"H{r}:J{r}"); merge(ws, f"K{r}:M{r}")
    merge(ws, "B76:D76"); merge(ws, "E76:G76"); merge(ws, "H76:J76"); merge(ws, "K76:M76")
    ws["B81"] = "Notes"; ws["B81"].font = hdr_font
    for r in range(73, 82):
        ws.row_dimensions[r].height = 13.35
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.print_area = "A1:M84"
    # conditional formatting ranges
    ws.conditional_formatting._cf_rules = {}
    log("Building Balance: 10 ERV rows added (supply on outside-air side, exhaust on exhaust side), totals re-based, building pressure table added (NEBB 8.19)")


# --------------------------------------------------------------------------- #
# ToC, Narrative sheet, sheet order
# --------------------------------------------------------------------------- #
def toc_and_narrative(wb):
    toc = wb["ToC"]
    moved = [m for m in toc.merged_cells.ranges if m.min_row >= 25]
    for m in moved:
        toc.unmerge_cells(str(m))
    toc.move_range("C25:L47", rows=3, cols=0)
    for m in moved:
        merge(toc, f"{get_column_letter(m.min_col)}{m.min_row+3}:{get_column_letter(m.max_col)}{m.max_row+3}")
    for r in range(47, 24, -1):
        toc.row_dimensions[r + 3].height = toc.row_dimensions[r].height
    for c in range(3, 13):
        copy_style(toc.cell(22, c), toc.cell(25, c)); copy_style(toc.cell(23, c), toc.cell(26, c))
    merge(toc, "C25:J26"); merge(toc, "K25:L26")
    toc["C25"] = "Energy Recovery Units - Data and Airflow"; toc["K25"] = "page 11"
    toc["C13"] = "Report Summary - Narrative and Remarks"
    toc["C22"] = "Make-up Air Units - Data, Airflow and Supply Methods"
    toc.print_area = "A1:M50"
    log("ToC: ERV entry inserted, summary/MAU entries renamed; run SyncToCPageCounts to refresh page numbers")
    # Narrative sheet from Summary - New
    src = wb["Summary - New"]
    ws = wb.copy_worksheet(src); ws.title = "Narrative"
    ws.sheet_properties.codeName = None
    ws.print_area = "A1:M41"; ws.sheet_properties.pageSetUpPr = copy.copy(src.sheet_properties.pageSetUpPr)
    ws.page_setup.fitToHeight = src.page_setup.fitToHeight; ws.page_setup.scale = src.page_setup.scale
    for m in [m for m in ws.merged_cells.ranges if m.min_row >= 12]:
        ws.unmerge_cells(str(m))
    for row in ws.iter_rows(min_row=12, max_row=62):
        for c in row:
            c.value = None; c.border = Border()
    ws["C5"] = "Report Summary - System Set-up Narrative"
    ws["C9"] = ("Describe the system set-up conditions established before testing, adjusting and balancing, and the rationale: how full-flow "
                "conditions were established, control configuration (BAS/DDC overrides, VFD commands, damper positions, filter condition), "
                "and the steps taken to reach the desired set-up. Deficiencies and items that could not be obtained are listed on the Remarks pages "
                "that follow, with the report page noted in the Comments column. (NEBB Procedural Standard 5.2.4)")
    ws["C9"].alignment = Alignment(wrap_text=True, vertical="top")
    merge(ws, "C12:L40")
    ws["C12"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["C12"].font = Font(name=src["C9"].font.name, size=10)
    for r in range(12, 41):
        for c in range(3, 13):
            ws.cell(r, c).border = Border(left=THIN if c == 3 else None, right=THIN if c == 12 else None,
                                          top=THIN if r == 12 else None, bottom=THIN if r == 40 else None)
    log("Narrative sheet added for the NEBB 5.2.4 system set-up narrative")


ORDER = ["Cover Page", "ToC", "Narrative", "Summary - New", "Summary - (E)", "{Project Information}", "{Equipment Data Entry}",
         "{Dropdowns}", "Building Balance", "RTU Data", "RTU Airflow", "MAU Data", "MAU Airflow", "MAU Supply Methods",
         "ERV Data", "ERV Airflow", "Fan Data (EFs, TFs, etc.)", "Fan Airflow", "VAV Data", "VAV 1-20 Airflow", "Hoods",
         "Traverses", "Photos", "Certification", "NEBB Cert ", "NEBB Frm Cert", "Abbreviations", "Calibration"]


def sheet_order(wb):
    wb._sheets = [wb[n] for n in ORDER] + [s for s in wb._sheets if s.title not in ORDER]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Cover Page"
    log("Sheet order set: " + ", ".join(s.title for s in wb.worksheets))


STEPS2 = [harmonize_formulas, dropdown_tables, data_entry_sections, erv_sheets, vav_links, hoods_rebuild, mau_methods,
          data_block_fields, airflow_sheet_fields, traverse_profiles, building_balance_erv, toc_and_narrative, sheet_order]
