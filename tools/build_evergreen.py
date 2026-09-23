"""Produce a cleaned reference copy of the CaptiveAire / Evergreen field worksheets.

    python3 tools/build_evergreen.py   -> "tb-worksheet-evergreen 3 (reference).xlsx"

Changes versus "tb-worksheet-evergreen 2.xlsx" (kept untouched as the backup):
  * all job-specific entries cleared (job number, hood/PSP dimensions, velocities, filter picks)
  * Supply Fan Filters: VLOOKUPs use exact match (FALSE) – approximate match silently
    returned the wrong row for an unlisted size
  * PSP Supply: CFM/ft no longer shows #DIV/0! when a block is empty
  * a README sheet documenting the constants and how the TAB workbook uses them
"""
import os
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "tb-worksheet-evergreen 2.xlsx")
OUT = os.path.join(ROOT, "tb-worksheet-evergreen 3 (reference).xlsx")

wb = load_workbook(SRC)
# --- clear entered data
for ws in wb.worksheets:
    ws["A2"] = "Job Number ="
ex = wb["Exhaust with Baffle Filters"]
for top in (6, 14, 22, 30):
    ex[f"D{top}"] = None; ex[f"E{top}"] = None; ex[f"B{top}"] = None
    for c in "BCDEFGHIJKL":
        ex[f"{c}{top+3}"] = "No Filter"; ex[f"{c}{top+4}"] = None
    ex[f"A{top+2}"] = "Baffle K Factor"
ps = wb["PSP Supply"]
for top in (6, 13, 20, 27):
    ps[f"B{top}"] = None; ps[f"B{top+1}"] = None; ps[f"B{top+2}"] = None
    for c in "BCDEFGHIJKLMNOPQ":
        ps[f"{c}{top+4}"] = None
    ps[f"S{top+4}"] = f'=IF(OR(B{top}="",R{top+4}=0),"",R{top+4}/(B{top}/12))'
    ps[f"R{top+4}"] = (f'=IF(OR(B{top}="",B{top+1}="",COUNT(B{top+4}:Q{top+4})=0),0,'
                       f'AVERAGE(B{top+4}:Q{top+4})*(B{top}-2-N(B{top+2})*2)*VLOOKUP(B{top+1},$V$7:$X$16,1,FALSE)*VLOOKUP(B{top+1},$V$7:$X$16,3,FALSE)/144)')
sf = wb["Supply Fan Filters"]
for top in (7, 14, 21, 28):
    for c in "BCDEFGHIJKLM":
        sf[f"{c}{top+1}"] = "No Filter"; sf[f"{c}{top+2}"] = None
        sf[f"{c}{top+3}"] = f"=VLOOKUP({c}{top+1},$Q$7:$S$17,2,FALSE)*N({c}{top+2})*VLOOKUP({c}{top+1},$Q$7:$S$17,3,FALSE)"
for name, sizes in (("HVC or Slot Filters", "$AL$7:$AL$9"), ("Condensate Baffle Filters", "$AL$7:$AL$15")):
    ws = wb[name]
    for top in (7, 14, 21, 28):
        for c in ("B", "E", "H", "K", "N", "Q", "T", "W", "Z", "AC", "AF"):
            ws[f"{c}{top+1}"] = "No Filter"
        for col in range(2, 35):
            ws.cell(top + 2, col).value = None
df = wb["Direct Fired Profile Pres. CFM"]
for r in range(4, 9):
    df[f"B{r}"] = 0
bp = wb["Building Pressures"]
bp["B6"] = None; bp["B7"] = None
# --- README
rd = wb.create_sheet("README", 0)
rd.column_dimensions["A"].width = 120
lines = [
    "CaptiveAire / Evergreen Telemetry field worksheets – cleaned reference copy",
    "",
    "Source: tb-worksheet-evergreen 2.xlsx (kept as backup). Job data removed; formulas unchanged except:",
    "  - Supply Fan Filters: VLOOKUP now uses exact match (FALSE).",
    "  - PSP Supply: CFM/ft returns blank instead of #DIV/0! when the block is empty; blanks count treated as 0.",
    "",
    "Constants used by the a2b TAB workbook ({Dropdowns} sheet, columns H:Z):",
    "  Exhaust baffle filters (VelGrid): CFM = velocity x free area x K, K = 1.34 Captrate / 1.28 baffle (rev 2.1)",
    "  Supply fan filters (VelGrid):     CFM = velocity x free area x 1.35 (rev 1)",
    "  Condensate baffle / HVC slot (Airfoil): CFM = average of 3 readings x probe constant (rev 0)",
    "  PSP supply (VelGrid): CFM = avg of readings x (length - 2 - 2 x blanks) x width x K / 144, K = 0.88 (<=12 in) or 0.95 (>=14 in) (rev 2)",
    "  Direct-fired burner profile pressure: CFM by housing size 1-5, linear interpolation between table points (rev 1)",
    "",
    "Note: the exhaust and supply tables list slightly different free areas for the same nominal filter size",
    "(e.g. 16x20: 1.73 vs 1.75 sq ft). Both are carried as published; confirm the current values with CaptiveAire.",
]
for i, t in enumerate(lines, 1):
    rd[f"A{i}"] = t
    rd[f"A{i}"].alignment = Alignment(wrap_text=True)
rd["A1"].font = Font(bold=True, size=12)
wb.save(OUT)
print("written", OUT)
